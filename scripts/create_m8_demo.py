"""Create a new, explicitly synthetic 15-order QA fixture; not research data."""

from datetime import datetime
from pathlib import Path
import sys

from openpyxl import load_workbook
from src.validate_sheet import SCHEMA, CODES
from src.pilot_report import REVIEW_HEADERS, _destination


def create_demo(path):
    path = _destination(path)
    root = Path(__file__).resolve().parents[1]
    w = load_workbook(root / "project_docs/Anchored_Returns_Pilot_Coding_Sheet_v0.1.2.xlsx")
    maps = {
        sheet: {v: i for i, v in enumerate(SCHEMA[key], 1)}
        for sheet, key in [("Cases", "case_headers"), ("Passages", "passage_headers")]
    }

    def put(sheet, row, values):
        for k, v in values.items():
            w[sheet].cell(row, maps[sheet][k]).value = v

    review = w.create_sheet("M3_Review")
    review.append(REVIEW_HEADERS)
    states = [
        "CT",
        "MA",
        "ME",
        "RI",
        "VT",
        "CT",
        "MA",
        "VT",
        "RI",
        "NH",
        "ME",
        "CT",
        "MA",
        "RI",
        "VT",
    ]
    relrow = 2
    for i, state in enumerate(states):
        cid = f"P{i + 1:02d}"
        year = 2019 + i % 3 if i < 5 else 2022 if i < 8 else 2023 + (i - 8) % 4
        if state == "NH":
            year = 2025
        c = dict(
            case_id=cid,
            sample_status="Included by design" if state == "NH" else "Sample",
            state=state,
            commission="Synthetic commission",
            docket="DE 24-070" if state == "NH" else f"TEST-{i + 1:02d}",
            utility="Synthetic utility " + str(i + 1),
            service="Electric",
            segment="Distribution",
            decision_type="Fully litigated",
            order_date=datetime(year, 6, 15),
            order_url="https://example.org/synthetic-test-only",
            roe_section_pages="1-3",
            requested_roe_initial=0.108,
            requested_roe=0.102,
            awarded_roe=round(0.09325 + i * 0.0002, 5),
            perf_adjustment_bp=-25 if i == 0 else None if i == 2 else 0,
            advocate_rec=None if i % 4 == 1 else 0.084 + i * 0.0001,
            staff_rec=0.09 + i * 0.00015,
            alj_rec=None if i % 3 else 0.092,
            nat_avg_benchmark=None if i % 5 == 0 else 0.095,
            prior_authorized_roe=0.092 + 0.0004 * (i % 3),
            prior_order_date=datetime(year - 4, 5, 1),
            ust10_at_prior_order=0.02 + 0.001 * (i % 4),
            ust10_at_order=0.01 + 0.002 * i,
            decisive_rationale="MKT" if i % 2 == 0 else "PEER",
            conclusory_flag="N",
            coder="SYNTHETIC_A",
            date_coded=datetime(2026, 9, 30),
            minutes_spent=30 + i * 3,
            sources_notes="SYNTHETIC TEST ONLY. All figures and classifications are fabricated software inputs, not docket facts. Test source p. 1.",
        )
        put("Cases", i + 3, c)
        for j in range(3):
            pid = f"{cid}-{j + 1:03d}"
            flags = {code: 0 for code in CODES}
            flags.update(
                MKT=int(j > 0 or i % 2 == 0),
                PEER=int(j == 0 and i % 2 == 1),
                OWN=int(j == 0 and i % 3 == 0),
                GRAD=int(j == 0 and i % 4 == 0),
                CONC=int(j == 0 and i % 5 == 0),
            )
            method = (
                "n.a."
                if flags["MKT"] == 0
                else "Expected earnings"
                if j == 1 and i % 3 == 0
                else "Multiple"
                if j == 1 and i % 4 == 0
                else "DCF"
            )
            put(
                "Passages",
                3 + i * 3 + j,
                dict(
                    passage_id=pid,
                    case_id=cid,
                    page_para="Synthetic p. 1",
                    section="Positions of parties" if j == 2 else "Commission analysis",
                    excerpt="SYNTHETIC TEST ONLY. This is not text from an order.",
                    method=method,
                    primary_code="MKT" if flags["MKT"] else "PEER",
                    stance="Summarized only" if j == 2 else "Relied on",
                    coder="SYNTHETIC_A",
                    **flags,
                ),
            )
            if j == 0:
                for code in ("MKT", "PEER", "OWN", "GRAD", "CONC"):
                    a = flags[code]
                    b = 1 - a if i == 1 and code in ("MKT", "PEER") else a
                    for col, v in enumerate((pid, code, a, b), 1):
                        w["Reliability"].cell(relrow, col, v)
                    relrow += 1
        if i < 12:
            review.append(
                [
                    cid,
                    int(i % 4 != 0),
                    "SYNTHETIC_REVIEWER",
                    "assistant",
                    "2026-09-30",
                    "Synthetic top-range comparison, not a source review.",
                ]
            )
    path.parent.mkdir(parents=True, exist_ok=True)
    w.save(path)
    w.close()
    return path


if __name__ == "__main__":
    print(create_demo(sys.argv[1]))
