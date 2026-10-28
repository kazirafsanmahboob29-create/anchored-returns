"""Read-only audit of the saved pilot inputs and M1–M8 outputs; print JSON.

Run from the repository root: python -m scripts.audit_project [--pdfinfo /path/to/pdfinfo]
This checks reproducibility and file integrity, not eligibility or research conclusions.
"""

import argparse
import csv
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
from importlib.metadata import version
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import zipfile

import pymupdf
from openpyxl import load_workbook

from src.locate_roe import locate_roe
from src.manifest import read_manifest
from src.segment import segment, HEADERS
from src.number_finder import find_numbers, format_table
from src.validate_sheet import validate_sheet, SCHEMA, _signature
from src.reliability import reliability, RELIABILITY_SCHEMA, ReliabilityMismatchError
from src.pilot_report import build_report

ROOT = Path(__file__).resolve().parents[1]


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads((ROOT / path).read_text())


def audit(pdfinfo=None):
    checked = {}
    originals = []
    with zipfile.ZipFile(ROOT.parent / "Energy Research Project.zip") as archive:
        for line in (ROOT / "project_docs/SHA256SUMS").read_text().splitlines():
            expected, name = line.split("  ", 1)
            path = ROOT / "project_docs" / name
            require(digest(path) == expected, f"Original changed: {name}")
            members = [m for m in archive.namelist() if Path(m).name == name]
            require(
                len(members) == 1 and archive.read(members[0]) == path.read_bytes(),
                f"Archive mismatch: {name}",
            )
            originals.append(name)
    require(
        (ROOT / "codebook/Anchored_Returns_Codebook_v0.1.2.docx").read_bytes()
        == (ROOT / "project_docs/Anchored_Returns_Codebook_v0.1.2.docx").read_bytes(),
        "Codebook copy changed",
    )
    checked["original_archive_members_unchanged"] = originals
    locked = [
        line.split("==")
        for line in (ROOT / "requirements-lock.txt").read_text().splitlines()
        if line.strip()
    ]
    for name, expected in locked:
        require(version(name) == expected, f"Installed package differs from lock: {name}")
    checked["locked_package_versions_checked"] = len(locked)

    sources = read_json("reports/source_inventory.json")["sources"]
    source_by_id = {s["source_id"]: s for s in sources}
    require(len(source_by_id) == len(sources), "Duplicate source IDs")
    pdf_checks = []
    for source in sources:
        path = ROOT / source["local_path"]
        require(digest(path) == source["sha256"], f"Source hash mismatch: {path}")
        require(path.stat().st_size == source["bytes"], f"Source size mismatch: {path}")
        with pymupdf.open(path) as doc:
            require(not doc.is_repaired and not doc.is_encrypted, f"Invalid source PDF: {path}")
            require(len(doc) == source["n_pages"], f"Source page count mismatch: {path}")
        record = dict(source_id=source["source_id"], n_pages=source["n_pages"])
        if pdfinfo:
            result = subprocess.run(
                [pdfinfo, str(path)], capture_output=True, text=True, check=True, timeout=30
            )
            match = re.search(r"^Pages:\s+(\d+)$", result.stdout, re.M)
            require(
                match and int(match[1]) == source["n_pages"], f"Poppler page-count mismatch: {path}"
            )
            record["poppler_page_count_agrees"] = True
        pdf_checks.append(record)
    checked["source_pdfs"] = pdf_checks
    registrations = []
    for name in ("data/manifest.csv", "data/practice_manifest.csv", "data/supporting_manifest.csv"):
        registrations.extend(read_manifest(ROOT / name))
    rows = {r["case_id"]: r for r in registrations}
    require(len(rows) == len(registrations), "Duplicate case ID across manifests")
    for case, row in rows.items():
        source = source_by_id[case]
        require(
            Path(row["file_path"]) == ROOT / source["local_path"], f"Manifest path mismatch: {case}"
        )
        for key in ("state", "docket", "sha256", "source_url", "download_date"):
            require(row[key] == source[key], f"Manifest/catalog mismatch: {case}/{key}")
        require(int(row["n_pages"]) == source["n_pages"], f"Manifest pages mismatch: {case}")
    checked["registered_documents_checked"] = len(rows)
    checked["document_roles"] = dict(
        candidate_orders=sum(k.startswith("CAND-") for k in rows),
        practice_orders=sum(k.startswith("PRACTICE-") for k in rows),
        supporting_documents=sum(k.startswith("SUPP-") for k in rows),
    )

    native_by_case, native_count = {}, 0
    for record in read_json("reports/m2_extraction_audit.json"):
        case = record["case_id"]
        require(record["source_sha256"] == rows[case]["sha256"], f"M2 source mismatch: {case}")
        require(
            digest(record["output_path"]) == record["output_sha256"],
            f"M2 output hash mismatch: {case}",
        )
        pages = read_json(record["output_path"])
        with pymupdf.open(rows[case]["file_path"]) as doc:
            require(len(pages) == len(doc), f"M2 page count mismatch: {case}")
            for number, (saved, page) in enumerate(zip(pages, doc), 1):
                text = page.get_text("text", sort=False)
                expected = dict(
                    page=number,
                    text=text,
                    n_chars=len(text),
                    likely_scanned=sum(not c.isspace() for c in text) < record["min_text_chars"],
                )
                require(saved == expected, f"Native extraction differs: {case} page {number}")
        require(
            [p["page"] for p in pages if p["likely_scanned"]] == record["flagged_pages"],
            f"Flag mismatch: {case}",
        )
        native_by_case[case] = pages
        native_count += len(pages)
    checked["native_pages_recomputed_and_compared"] = native_count
    ocr = read_json("reports/m2_ocr_audit.json")
    require(digest(ocr["output_path"]) == ocr["output_sha256"], "OCR output hash mismatch")
    recognized = read_json(ocr["output_path"])
    native = native_by_case[ocr["case_id"]]
    require(len(recognized) == len(native), "OCR page count mismatch")
    require(ocr["ocr_pages"] == ocr["flagged_pages"], "OCR processed nonflagged pages")
    for old, new in zip(native, recognized):
        require(
            set(new) == set(old)
            and new["page"] == old["page"]
            and new["n_chars"] == len(new["text"])
            and new["likely_scanned"] == old["likely_scanned"],
            "OCR schema/flag mismatch",
        )
        if not old["likely_scanned"]:
            require(old == new, "OCR changed an unflagged page")
    for dependency in read_json("reports/ocr_dependency.json")["files"]:
        require(
            digest(ROOT / dependency["local_path"]) == dependency["sha256"],
            "OCR dependency hash mismatch",
        )
    checked["ocr_audit"] = dict(
        pages=len(recognized),
        processed_pages=len(ocr["ocr_pages"]),
        unflagged_pages_unchanged=True,
        recognition_accuracy="not certified",
    )

    fred = read_json("reports/fred_verification.json")
    for role in ("raw", "normalized"):
        require(
            digest(ROOT / fred[role + "_path"]) == fred[role + "_sha256"],
            f"FRED {role} hash mismatch",
        )
    with (ROOT / fred["raw_path"]).open(newline="") as f:
        raw = list(csv.DictReader(f))
    with (ROOT / fred["normalized_path"]).open(newline="") as f:
        normalized = list(csv.DictReader(f))
    require(raw and set(raw[0]) == {"observation_date", "DGS10"}, "FRED raw schema changed")
    require(
        normalized and set(normalized[0]) == {"date", "ust10", "source_url"},
        "FRED normalized schema changed",
    )
    require(len(raw) == len(normalized) == fred["rows"], "FRED row count mismatch")
    dates = [r["observation_date"] for r in raw]
    require(dates == sorted(set(dates)), "FRED dates not unique and ordered")
    missing = 0
    for old, new in zip(raw, normalized):
        require(new["source_url"] == fred["series_url"], "FRED row provenance changed")
        require(old["observation_date"] == new["date"], "FRED date shifted")
        if old["DGS10"] in ("", ".", "NA"):
            missing += 1
            require(new["ust10"] == "", "FRED missing value filled")
        else:
            require(
                Decimal(old["DGS10"]) == 100 * Decimal(new["ust10"]),
                "FRED unit conversion mismatch",
            )
    require(missing == fred["missing"], "FRED missing count mismatch")
    checked["fred_rows_checked"] = dict(total=len(raw), missing=missing)
    m3 = []
    for path in sorted((ROOT / "reports/m3_suggestions").glob("*.json")):
        saved = read_json(path)
        config = saved["config"]
        actual = locate_roe(
            saved["source_path"],
            **{k: config[k] for k in ("top_n", "max_gap_pages", "min_page_score")},
        )
        require(saved == actual, f"M3 output no longer reproduces: {path.name}")
        m3.append(path.name)
    require(len(m3) == len(rows), "Missing M3 registered-document output")
    checked["m3_outputs_reproduced"] = m3
    summary = read_json("reports/research_summary_provenance.json")
    require(digest(ROOT / summary["archived_path"]) == summary["sha256"], "Project summary changed")
    checked["owner_summary_preserved"] = True
    m4 = []
    extraction_by_id = {r["case_id"]: r for r in read_json("reports/m2_extraction_audit.json")}
    template = ROOT / "project_docs/Anchored_Returns_Pilot_Coding_Sheet_v0.1.2.xlsx"
    original = load_workbook(template)
    for plan_path in sorted((ROOT / "reports/m4_review_plans").glob("*.json")):
        plan = read_json(plan_path)
        case = plan["case_id"]
        bundle = ROOT / "data/drafts" / (case + "-v1")
        saved = read_json(bundle / "audit.json")
        actual = segment(extraction_by_id[case]["output_path"], rows[case]["file_path"], plan_path)
        require(all(saved[k] == v for k, v in actual.items()), f"M4 no longer reproduces: {case}")
        require(
            read_json(bundle / "passages.json") == actual["passages"],
            f"M4 passage JSON differs: {case}",
        )
        for filename, key in [
            ("passages.json", "passages_sha256"),
            (saved["workbook"], "workbook_sha256"),
        ]:
            require(
                digest(bundle / filename) == saved[key], f"M4 artifact changed: {case}/{filename}"
            )
        require(digest(template) == saved["template_sha256"], "M4 template changed")
        workbook = load_workbook(bundle / saved["workbook"])
        for name in original.sheetnames:
            require(
                str(original[name].data_validations) == str(workbook[name].data_validations),
                f"Validation changed: {name}",
            )
            for row in original[name]:
                for cell in row:
                    if name != "Passages" or cell.row < 3:
                        require(
                            workbook[name][cell.coordinate].value == cell.value,
                            f"Template value changed: {name}/{cell.coordinate}",
                        )
        for number, passage in enumerate(actual["passages"], 3):
            for col, field in enumerate(HEADERS, 1):
                cell = workbook["Passages"].cell(number, col)
                expected = (
                    passage.get(field)
                    if field in ("passage_id", "case_id", "page_para", "excerpt", "notes")
                    else None
                )
                require(
                    cell.value == expected and cell.data_type != "f",
                    f"M4 entry/label mismatch: {number}/{field}",
                )
        m4.append(
            dict(
                case_id=case,
                passages=len(actual["passages"]),
                source_lines=len(actual["source_line_ledger"]),
                reviewer_type=plan["reviewer_type"],
                human_validation_complete=False,
            )
        )
    require(bool(m4), "Missing M4 practice output")
    checked["m4_outputs_reproduced"] = m4
    m5 = {}
    for record in read_json("reports/m5_extraction_audit.json"):
        case = record["case_id"]
        require(case in rows and case not in m5, f"Unknown/duplicate M5 document: {case}")
        expected_text = (
            ocr["output_path"] if case == ocr["case_id"] else extraction_by_id[case]["output_path"]
        )
        require(Path(record["source_path"]) == Path(expected_text), f"M5 input mismatch: {case}")
        actual = find_numbers(record["source_path"], **record["config"])
        require(actual["source_sha256"] == record["source_sha256"], f"M5 input changed: {case}")
        require(read_json(record["output_path"]) == actual, f"M5 no longer reproduces: {case}")
        require(
            digest(ROOT / record["output_path"]) == record["output_sha256"],
            f"M5 JSON changed: {case}",
        )
        require(
            (ROOT / record["table_path"]).read_text() == format_table(actual),
            f"M5 table differs: {case}",
        )
        require(
            digest(ROOT / record["table_path"]) == record["table_sha256"],
            f"M5 table changed: {case}",
        )
        require(actual["n_candidates"] == record["n_candidates"], f"M5 count differs: {case}")
        m5[case] = actual
    require(set(m5) == set(rows), "Missing M5 registered-document output")
    checked["m5_outputs_reproduced"] = {k: v["n_candidates"] for k, v in m5.items()}
    ledger = read_json("data/research/assistant_verified_case_facts.json")
    facts = ledger["facts"]
    require(
        len(facts) == ledger["n_facts"] == len({r["fact_id"] for r in facts}),
        "Fact IDs/count mismatch",
    )
    for fact in facts:
        source = source_by_id[fact["source_id"]]
        require(
            fact["pdf_sha256"] == source["sha256"]
            and fact["source_url"] == source["source_url"]
            and fact["pdf_path"] == source["local_path"]
            and fact["event_date"] == source["order_date"],
            f"Fact source metadata mismatch: {fact['fact_id']}",
        )
        require(digest(ROOT / fact["text_path"]) == fact["text_sha256"], "Fact text changed")
        text = read_json(fact["text_path"])[fact["pdf_page"] - 1]["text"]
        evidence = fact["evidence"]
        require(
            text[evidence["start_char"] : evidence["end_char"]] == evidence["exact_text"],
            "Fact evidence mismatch",
        )
        candidates = {c["candidate_id"]: c for c in m5[fact["source_id"]]["candidates"]}
        candidate = candidates[fact["candidate_id"]]
        require(
            candidate["page"] == fact["pdf_page"]
            and candidate["value"] == fact["printed_value"]
            and candidate["unit"] == fact["printed_unit"]
            and candidate["fraction"] == fact["value_fraction"],
            "Fact/candidate provenance mismatch",
        )
        require(
            Decimal(fact["printed_value"])
            == Decimal(fact["value_fraction"])
            * (10000 if fact["printed_unit"] == "basis_points" else 100),
            "Fact unit conversion mismatch",
        )
        require(fact["verification"] == "assistant_source_checked", "Verification type changed")
        require(
            f"PDF viewer p. {fact['pdf_page']}" in fact["sources_notes"]
            and fact["source_url"] in fact["sources_notes"],
            "Missing fact source note/page",
        )
    ct = {
        r["field"]: Decimal(r["value_fraction"])
        for r in facts
        if r["case_id"] == "CAND-CT-22-08-08"
    }
    require(
        ct["awarded_roe_before_performance_reductions"] - ct["performance_reduction_magnitude"]
        == ct["awarded_roe_after_performance_reductions"],
        "CT base/reduction/final reconciliation failed",
    )
    checked["assistant_fact_provenance_checked"] = len(facts)
    checked["ct_base_minus_reduction_equals_final"] = True
    require(SCHEMA["template_sha256"] == digest(template), "M6 schema/template identity changed")
    require(
        SCHEMA["case_headers"] == [c.value for c in original["Cases"][1]], "M6 Cases schema changed"
    )
    require(
        SCHEMA["passage_headers"] == [c.value for c in original["Passages"][1]],
        "M6 Passages schema changed",
    )
    require(
        SCHEMA["reliability_headers"] == [c.value for c in original["Reliability"][1][:4]],
        "M6 Reliability schema changed",
    )
    require(
        SCHEMA["lists"]
        == {
            col[0].value: [c.value for c in col[1:] if c.value is not None]
            for col in original["Lists"].iter_cols()
        },
        "M6 dropdown schema changed",
    )
    require(
        SCHEMA["case_formulas"] == {c.coordinate: c.value for c in original["Cases"][2][41:]},
        "M6 formula schema changed",
    )
    for sheet, width in [("Cases", 41), ("Passages", 21)]:
        require(
            SCHEMA["example_signatures"][sheet] == _signature(list(original[sheet][2][:width])),
            "M6 example identity changed",
        )
    m6 = read_json("reports/m6_validation_audit.json")
    runs = []
    for run in m6["runs"]:
        require(digest(ROOT / run["source_path"]) == run["source_sha256"], "M6 workbook changed")
        require(digest(ROOT / run["result_path"]) == run["result_sha256"], "M6 report changed")
        actual = validate_sheet(ROOT / run["source_path"], **run["config"])
        require(actual == read_json(run["result_path"]), f"M6 no longer reproduces: {run['name']}")
        if run["expected_error"]:
            require(
                any(
                    i["code"] == run["expected_error"] and i["severity"] == "error"
                    for i in actual["issues"]
                ),
                "M6 missed deliberate/known defect",
            )
        else:
            require(actual["errors"] == 0, f"Unexpected M6 defect: {run['name']}")
        runs.append(
            dict(
                name=run["name"],
                status=actual["status"],
                errors=actual["errors"],
                incomplete=actual["incomplete"],
            )
        )
    patch = read_json(m6["practice_patch"])
    require(
        digest(ROOT / patch["source_path"]) == patch["source_sha256"], "M6 practice source changed"
    )
    require(
        digest(ROOT / patch["output_path"]) == patch["output_sha256"], "M6 working copy changed"
    )
    require(
        digest(ROOT / patch["source_pdf"]) == patch["source_pdf_sha256"],
        "Practice metadata source changed",
    )
    before = load_workbook(ROOT / patch["source_path"])
    after = load_workbook(ROOT / patch["output_path"])
    changes = {(x["sheet"], x["cell"]): x for x in patch["changes"]}
    require(before.sheetnames == after.sheetnames, "M6 changed workbook tabs")
    for sheet in before.sheetnames:
        a, b = before[sheet], after[sheet]
        require(
            a.sheet_state == b.sheet_state and a.freeze_panes == b.freeze_panes,
            "M6 changed tab visibility/freeze panes",
        )
        require(str(a.data_validations) == str(b.data_validations), "M6 changed dropdowns")
        require(
            (a.max_row, a.max_column) == (b.max_row, b.max_column), "M6 changed sheet dimensions"
        )
        for row in a:
            for cell in row:
                new = b[cell.coordinate]
                change = changes.get((sheet, cell.coordinate))
                if change:
                    value = new.value.isoformat() if isinstance(new.value, datetime) else new.value
                    require(
                        cell.value == change["before"] and value == change["after"],
                        "M6 metadata patch mismatch",
                    )
                else:
                    require(
                        cell.value == new.value and cell.data_type == new.data_type,
                        f"M6 changed preserved data: {sheet}!{cell.coordinate}",
                    )
                require(cell._style == new._style, "M6 changed cell formatting")
                require(
                    (cell.comment.text if cell.comment else None)
                    == (new.comment.text if new.comment else None),
                    "M6 changed cell notes",
                )
        for key in set(a.column_dimensions) | set(b.column_dimensions):
            da, db = a.column_dimensions[key], b.column_dimensions[key]
            expected = 30 if sheet == "Cases" and key == "A" else da.width
            require(
                db.width == expected and da.hidden == db.hidden,
                "M6 changed column layout unexpectedly",
            )
        for key in set(a.row_dimensions) | set(b.row_dimensions):
            da, db = a.row_dimensions[key], b.row_dimensions[key]
            expected = 150 if sheet == "Cases" and key == 3 else da.height
            require(
                db.height == expected and da.hidden == db.hidden,
                "M6 changed row layout unexpectedly",
            )
    before.close()
    after.close()
    checked["m6_outputs_reproduced"] = runs
    checked["m6_practice_copy_preserves_original_except_logged_metadata_and_layout"] = True
    template_path = ROOT / "project_docs/Anchored_Returns_Pilot_Coding_Sheet_v0.1.2.xlsx"
    require(
        digest(template_path) == RELIABILITY_SCHEMA["template_sha256"],
        "M7 template identity changed",
    )
    template = load_workbook(template_path)
    sheet = template["Reliability"]
    require(
        [c.value for c in sheet[1]] == RELIABILITY_SCHEMA["headers"],
        "M7 header schema differs from template",
    )
    require(
        [sheet.cell(r, 6).value for r in range(2, 13)] == RELIABILITY_SCHEMA["codes"],
        "M7 code order differs",
    )
    require(
        {sheet.cell(2, c).coordinate: sheet.cell(2, c).value for c in range(7, 16)}
        == RELIABILITY_SCHEMA["formulas"],
        "M7 formula schema differs",
    )
    template.close()
    m7 = read_json("reports/m7_validation_audit.json")
    require(
        m7["synthetic_ratings_are_research_data"] is False,
        "M7 fixtures must not become research ratings",
    )
    runs = []
    for run in m7["runs"]:
        require(digest(ROOT / run["source_path"]) == run["source_sha256"], "M7 input changed")
        require(digest(ROOT / run["result_path"]) == run["result_sha256"], "M7 report changed")
        raised = None
        try:
            actual = reliability(ROOT / run["source_path"], **run["config"])
        except ReliabilityMismatchError as error:
            actual = error.report
            raised = type(error).__name__
        require(raised == run["expected_exception"], f"M7 exception differs: {run['name']}")
        require(actual == read_json(run["result_path"]), f"M7 output differs: {run['name']}")
        runs.append(
            dict(name=run["name"], status=actual["status"], total_pairs=actual["total_pairs"])
        )
    evidence = m7["engine_evidence"]
    require(digest(ROOT / evidence["path"]) == evidence["sha256"], "M7 engine evidence changed")
    engine = read_json(evidence["path"])
    require(engine["syntheticOnly"] and not engine["nativeExcelTested"], "M7 engine scope changed")
    for stage in engine["evidence"]:
        result = read_json(f"reports/m7_reliability/engine-{stage['name']}.json")
        require(result["status"] == "passed", "M7 engine output did not cross-check")
        require(
            result["codes"][0]["kappa"] == stage["expectedKappa"] == stage["summary"][7],
            "M7 recalculation result differs",
        )
    for visual in m7["visual_checks"]:
        require(digest(ROOT / visual["path"]) == visual["sha256"], "M7 visual evidence changed")
    checked["m7_outputs_reproduced"] = runs
    checked["m7_engine_change_restore_verified"] = True
    m8 = read_json("reports/m8_validation_audit.json")
    require(m8["synthetic_ratings_are_research_data"] is False, "M8 fixtures must remain synthetic")
    runs = []
    for run in m8["runs"]:
        output = ROOT / run["output_dir"]
        require(digest(ROOT / run["source_path"]) == run["source_sha256"], "M8 input changed")
        require(
            digest(output / "audit.json") == run["audit_sha256"], "M8 completion record changed"
        )
        saved = read_json(str(output / "audit.json"))
        require(
            sorted(str(p.relative_to(output)) for p in output.rglob("*") if p.is_file())
            == sorted([f["path"] for f in saved["files"]] + ["audit.json"]),
            "M8 bundle has missing or unexpected files",
        )
        for f in saved["files"]:
            require(digest(output / f["path"]) == f["sha256"], "M8 output changed: " + f["path"])
        with tempfile.TemporaryDirectory(prefix="anchored-m8-audit-") as temporary:
            regenerated = build_report(
                ROOT / run["source_path"], Path(temporary).resolve() / "report", **run["config"]
            )
            require(regenerated == saved, "M8 figures/tables/memo did not reproduce exactly")
        with pymupdf.open(output / "memo.pdf") as pdf:
            require(len(pdf) == 2, "M8 memo must have two pages")
        runs.append(
            dict(
                name=run["name"],
                files=len(saved["files"]) + 1,
                n_cases=saved["n_cases"],
                n_scatter=saved["n_scatter"],
            )
        )
    for image in m8["visual_checks"]:
        require(digest(ROOT / image["path"]) == image["sha256"], "M8 PDF visual evidence changed")
    checked["m8_bundles_reproduced_byte_for_byte"] = runs
    return dict(
        checked_at_utc=datetime.now(timezone.utc).isoformat(),
        status="passed",
        checks=checked,
        limits=[
            "No sample eligibility certification or human hit-rate measurement.",
            "No full audit of original workbook formulas or source-paper statistical methods.",
            "No fresh downloads or certification of OCR recognition accuracy.",
        ],
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdfinfo", help="Optional Poppler binary for independent PDF page counts")
    try:
        result = audit(**vars(parser.parse_args()))
    except (ValueError, OSError, KeyError, subprocess.SubprocessError) as error:
        print(f"audit failed: {error}", file=sys.stderr)
        return 2
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
