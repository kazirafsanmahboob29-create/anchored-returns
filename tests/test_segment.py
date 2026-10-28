"""M4 behavioral checks use synthetic text; research labels are never generated."""

import json
from pathlib import Path

from openpyxl import load_workbook
import pymupdf
import pytest

from src.segment import SegmentationError, digest, segment, write_draft

TEMPLATE = (
    Path(__file__).resolve().parents[1]
    / "project_docs/Anchored_Returns_Pilot_Coding_Sheet_v0.1.2.xlsx"
)


@pytest.fixture
def inputs(tmp_path):
    pdf = tmp_path / "order.pdf"
    with pymupdf.open() as doc:
        p = doc.new_page()
        for xy, text in [
            ((72, 40), "RUNNING HEADER"),
            ((72, 110), "ROE section"),
            ((108, 140), "A regu-"),
            ((72, 165), "lation with cost-of-"),
            ((72, 190), "capital continues"),
            ((72, 650), "1"),
            ((108, 650), "A footnote continuing"),
        ]:
            p.insert_text(xy, text)
        p = doc.new_page()
        for xy, text in [
            ((72, 40), "RUNNING HEADER"),
            ((72, 110), "across pages."),
            ((108, 140), "=not a formula"),
            ((72, 165), "RUNNING HEADER"),
            ((72, 195), "NEXT SECTION"),
            ((72, 225), "Excluded text."),
            ((108, 620), "on the second page."),
            ((72, 650), "2"),
            ((108, 650), "Second note."),
        ]:
            p.insert_text(xy, text)
        doc.save(pdf)
        pages = [
            dict(page=i + 1, text=p.get_text(), n_chars=len(p.get_text()), likely_scanned=False)
            for i, p in enumerate(doc)
        ]
    text = tmp_path / "text.json"
    text.write_text(json.dumps(pages))
    continuation = next(
        i for i, s in enumerate(pages[1]["text"].splitlines(), 1) if s == "on the second page."
    )
    plan = dict(
        case_id="TEST-M4",
        start_page=1,
        end_page=2,
        reviewer="Synthetic reviewer",
        reviewer_type="assistant",
        text_sha256=digest(text),
        pdf_sha256=digest(pdf),
        layout=dict(
            body_left=72,
            paragraph_indent=36,
            header_bottom=60,
            footer_top=760,
            header_patterns=["RUNNING HEADER"],
        ),
        start_at="ROE section",
        end_before="NEXT SECTION",
        dehyphenate_words=["regulation"],
        footnote_continuations=[f"2:{continuation}"],
    )
    path = tmp_path / "plan.json"
    path.write_text(json.dumps(plan))
    return dict(text_json=text, pdf=pdf, review_plan=path), plan


def alter(inputs, **changes):
    args, plan = inputs
    plan.update(changes)
    args["review_plan"].write_text(json.dumps(plan))
    return args


def test_cross_page_paragraph_and_separate_continued_notes(inputs):
    result = segment(**inputs[0])
    rows = result["passages"]
    joined = " ".join(r["excerpt"] for r in rows if r["kind"] == "body")
    assert "regulation with cost-of-capital continues across pages." in joined
    assert "Excluded text" not in joined
    assert "RUNNING HEADER" in joined  # Same words inside body must survive.
    notes = [r for r in rows if r["kind"] == "footnote"]
    assert [r["excerpt"] for r in notes] == [
        "1 A footnote continuing on the second page.",
        "2 Second note.",
    ]
    assert len(set(notes[0]["source_line_ids"])) == 3
    assert any("keep_hyphen" in t["operation"] for r in rows for t in r["normalization"])


def test_every_nonblank_line_has_one_disposition(inputs):
    result = segment(**inputs[0])
    ids = [s for r in result["passages"] for s in r["source_line_ids"]]
    assert len(ids) == len(set(ids))
    ledger = result["source_line_ledger"]
    assert all(l["disposition"] != "pending" for l in ledger)
    assert set(ids) == {l["id"] for l in ledger if l["disposition"].startswith("TEST-M4-")}


@pytest.mark.parametrize(
    "changes",
    [
        {"start_page": True},
        {"end_page": 3},
        {"start_page": 0},
        {"end_page": 0},
        {"reviewer_type": "verified"},
        {"reviewer": ""},
        {"case_id": "../bad"},
        {"text_sha256": "bad"},
        {"pdf_sha256": "bad"},
        {"start_at": "missing"},
        {"end_before": "missing"},
        {"paragraph_breaks": ["1:999"]},
    ],
)
def test_rejects_unreviewed_or_invalid_inputs(inputs, changes):
    with pytest.raises((SegmentationError, ValueError)):
        segment(**alter(inputs, **changes))


def test_scanned_or_ocr_altered_input_rejected(inputs):
    args, plan = inputs
    rows = json.loads(args["text_json"].read_text())
    rows[0]["likely_scanned"] = True
    args["text_json"].write_text(json.dumps(rows))
    with pytest.raises(SegmentationError, match="non-scanned"):
        segment(**alter(inputs, text_sha256=digest(args["text_json"])))


def test_same_hash_does_not_allow_pdf_text_mismatch(inputs):
    args, plan = inputs
    rows = json.loads(args["text_json"].read_text())
    rows[0]["text"] = rows[0]["text"].replace("regu-", "fake-")
    args["text_json"].write_text(json.dumps(rows))
    with pytest.raises(SegmentationError, match="native"):
        segment(**alter(inputs, text_sha256=digest(args["text_json"])))


def test_export_preserves_template_formulas_validations_and_blank_labels(inputs, tmp_path):
    args, _ = inputs
    before = {p: p.read_bytes() for p in [*args.values(), TEMPLATE]}
    out = tmp_path / "draft"
    write_draft(**args, template=TEMPLATE, output_dir=out)
    for p, content in before.items():
        assert p.read_bytes() == content
    original = load_workbook(TEMPLATE)
    actual = load_workbook(out / "passages_draft.xlsx")
    rows = json.loads((out / "passages.json").read_text())
    for name in original.sheetnames:
        for row in original[name]:
            for cell in row:
                if name != "Passages" or cell.row < 3:
                    assert actual[name][cell.coordinate].value == cell.value
        assert str(actual[name].data_validations) == str(original[name].data_validations)
    for n, row in enumerate(rows, 3):
        assert actual["Passages"].cell(n, 5).value == row["excerpt"]
        assert actual["Passages"].cell(n, 5).data_type == "s"
        for col in [4, *range(6, 21)]:
            assert actual["Passages"].cell(n, col).value is None
    assert actual["Summary"].sheet_state == "hidden"
    assert actual["Cases"].column_dimensions["AP"].hidden
    assert actual["Passages"]["A2"].value == original["Passages"]["A2"].value
    assert json.loads((out / "audit.json").read_text())["workbook_sha256"] == digest(
        out / "passages_draft.xlsx"
    )


@pytest.mark.parametrize("mode", ["exists", "protected", "symlink", "populated", "bad_schema"])
def test_no_overwrite_or_live_workbook_write(inputs, tmp_path, mode):
    args, _ = inputs
    template = TEMPLATE
    out = tmp_path / "draft"
    if mode == "exists":
        out.mkdir()
        (out / "keep.txt").write_text("keep")
    if mode == "protected":
        out = tmp_path / "coding" / "draft"
    if mode == "symlink":
        (tmp_path / "link").symlink_to(tmp_path, target_is_directory=True)
        out = tmp_path / "link" / "draft"
    if mode in ["populated", "bad_schema"]:
        wb = load_workbook(TEMPLATE)
        wb["Passages"]["A3" if mode == "populated" else "A1"] = "already filled"
        template = tmp_path / "template.xlsx"
        wb.save(template)
    with pytest.raises(SegmentationError):
        write_draft(**args, template=template, output_dir=out)
    if mode == "exists":
        assert (out / "keep.txt").read_text() == "keep"
    else:
        assert not out.exists()
