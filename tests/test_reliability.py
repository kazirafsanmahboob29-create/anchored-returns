"""M7 contract tests. All ratings and cached results below are synthetic."""

from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
import pytest

from src.reliability import (
    binary_agreement,
    reliability,
    ReliabilityInputError,
    ReliabilityMismatchError,
)

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "project_docs/Anchored_Returns_Pilot_Coding_Sheet_v0.1.2.xlsx"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
TOY_A = [1, 1, 1, 1, 0, 0, 0, 0]
TOY_B = [1, 1, 1, 0, 1, 0, 0, 0]
# Hand calculation: 6/8 agree; expected agreement 1/2; kappa = 1/2.
TOY_CACHE = dict(G2=8, H2=3, I2=3, J2=4, K2=4, L2=0.75, M2=0.5, N2=0.5, O2="Rewrite")


def inject_test_caches(path, values):
    """Write explicitly simulated caches, not evidence of Excel recalculation."""
    raw = path.read_bytes()
    with ZipFile(BytesIO(raw)) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet5.xml":
                root = ET.fromstring(content)
                for c in root.iter(NS + "c"):
                    if c.attrib["r"] not in values:
                        continue
                    value = values[c.attrib["r"]]
                    v = c.find(NS + "v")
                    if v is not None:
                        c.remove(v)
                    c.attrib.pop("t", None)
                    if value is not None:
                        if isinstance(value, str):
                            c.set("t", "e" if value.startswith("#") else "str")
                        ET.SubElement(c, NS + "v").text = str(value)
                content = ET.tostring(root)
            target.writestr(item, content)


def make_book(path, a=TOY_A, b=TOY_B, caches=TOY_CACHE):
    w = load_workbook(TEMPLATE)
    w["Cases"]["D3"] = "SYNTHETIC M7 SOFTWARE TEST"
    for i, (left, right) in enumerate(zip(a, b), 2):
        pid = f"P01-{i - 1:03d}"
        for col, value in enumerate((pid, "MKT", left, right), 1):
            w["Reliability"].cell(i, col, value)
        for col, value in enumerate(
            (pid, "P01", "Synthetic p. 1", None, "SYNTHETIC TEST ONLY. Not an order passage."), 1
        ):
            w["Passages"].cell(i + 1, col, value)
    w.save(path)
    w.close()
    if caches:
        inject_test_caches(path, caches)
    return path


def change(path, sheet, address, value):
    w = load_workbook(path)
    w[sheet][address] = value
    w.save(path)
    w.close()


@pytest.mark.parametrize(
    "a,b,expected,observed",
    [
        (TOY_A, TOY_B, 0.5, 0.75),
        ([0, 1], [0, 1], 1, 1),
        ([0, 1], [1, 0], -1, 0),
        ([0, 0, 1, 1], [0, 1, 0, 1], 0, 0.5),
        ([0, 0], [1, 1], 0, 0),
        ([1, 1, 1, 0, 0], [1, 1, 0, 0, 0], 8 / 13, 0.8),
    ],
)
def test_known_results_and_symmetry(a, b, expected, observed):
    result = binary_agreement(a, b)
    assert result["kappa"] == pytest.approx(expected)
    assert result["observed_agreement"] == observed
    assert binary_agreement(b, a)["kappa"] == pytest.approx(expected)


@pytest.mark.parametrize("values", [[0, 0], [1, 1], [1]])
def test_constant_agreement_is_undefined(values):
    r = binary_agreement(values, values)
    assert r["kappa"] is None and r["status"] == "undefined"
    assert r["observed_agreement"] == 1


def test_no_pairs_is_not_perfect_agreement():
    r = binary_agreement([], [])
    assert r["status"] == "no_pairs" and r["kappa"] is None


@pytest.mark.parametrize(
    "a,b",
    [([1], []), ([True], [1]), (["0"], [0]), ([None], [0]), ([2], [0]), ([float("nan")], [0])],
)
def test_bad_vectors_rejected(a, b):
    with pytest.raises(ReliabilityInputError):
        binary_agreement(a, b)


def test_toy_matches_saved_kappa_read_only(tmp_path):
    p = make_book(tmp_path / "toy.xlsx")
    before = p.read_bytes()
    r = reliability(p)
    assert p.read_bytes() == before
    assert r["status"] == "passed" and r["compared_codes"] == 1
    m = r["codes"][0]
    assert m["kappa"] == 0.5 and m["absolute_difference"] == 0
    assert m["comparison"] == "matched" and m["n_pairs"] == 8
    assert len(r["codes"]) == 11
    assert r["codes"][1]["status"] == "no_pairs"
    json.dumps(r, allow_nan=False)


@pytest.mark.parametrize(
    "cache,should_raise",
    [(0.499, False), (0.501, False), (0.5009999, False), (0.5010001, True), (0.8, True)],
)
def test_tolerance_boundary(tmp_path, cache, should_raise):
    p = make_book(tmp_path / "toy.xlsx", caches={**TOY_CACHE, "N2": cache})
    if should_raise:
        with pytest.raises(ReliabilityMismatchError) as error:
            reliability(p)
        assert error.value.report["status"] == "failed"
        assert "kappa_mismatch" in {i["code"] for i in error.value.report["issues"]}
    else:
        assert reliability(p)["status"] == "passed"


@pytest.mark.parametrize(
    "cell,value",
    [
        ("G2", 7),
        ("H2", 2),
        ("I2", 5),
        ("J2", 3),
        ("K2", 5),
        ("L2", 0.7),
        ("M2", 0.6),
        ("O2", "Yes"),
        ("N2", "#DIV/0!"),
        ("N2", "0.5"),
    ],
)
def test_stale_or_invalid_saved_summaries_raise(tmp_path, cell, value):
    p = make_book(tmp_path / "toy.xlsx", caches={**TOY_CACHE, cell: value})
    with pytest.raises(ReliabilityMismatchError):
        reliability(p)


def test_missing_cache_is_unverified_and_compute_only_is_explicit(tmp_path):
    p = make_book(tmp_path / "toy.xlsx", caches=None)
    r = reliability(p)
    assert r["status"] == "incomplete" and r["codes"][0]["comparison"] == "missing_cache"
    r = reliability(p, compute_only=True)
    assert r["status"] == "computed_only" and r["compared_codes"] == 0
    assert r["codes"][0]["kappa"] == 0.5 and r["codes"][0]["comparison"] == "not_requested"


def test_empty_template_returns_no_data():
    r = reliability(TEMPLATE)
    assert r["status"] == "no_data" and r["total_pairs"] == 0
    assert r["compared_codes"] == 0


@pytest.mark.parametrize("value", [0, 1])
def test_undefined_workbook_is_never_a_success_or_zero(tmp_path, value):
    caches = dict(
        G2=2,
        H2=2 * value,
        I2=2 * (1 - value),
        J2=2 * value,
        K2=2 * value,
        L2=1,
        M2=1,
        N2=None,
        O2=None,
    )
    p = make_book(tmp_path / "constant.xlsx", [value] * 2, [value] * 2, caches)
    r = reliability(p)
    assert r["status"] == "incomplete" and r["codes"][0]["comparison"] == "undefined"
    assert r["codes"][0]["kappa"] is None
    inject_test_caches(p, {"N2": 1})
    with pytest.raises(ReliabilityMismatchError):
        reliability(p)


def test_no_data_cannot_hide_a_stale_numeric_kappa(tmp_path):
    p = make_book(tmp_path / "empty.xlsx", [], [], {"N2": 0.8})
    with pytest.raises(ReliabilityMismatchError):
        reliability(p)


def test_one_missing_cache_does_not_hide_another_mismatch(tmp_path):
    p = make_book(tmp_path / "toy.xlsx", caches={**TOY_CACHE, "G2": None, "N2": 0.9})
    with pytest.raises(ReliabilityMismatchError):
        reliability(p)


def test_date_formatted_cache_has_json_readable_diagnostic(tmp_path):
    p = make_book(tmp_path / "toy.xlsx", caches=None)
    w = load_workbook(p)
    w["Reliability"]["N2"].number_format = "yyyy-mm-dd"
    w.save(p)
    w.close()
    inject_test_caches(p, TOY_CACHE)
    with pytest.raises(ReliabilityMismatchError) as error:
        reliability(p)
    json.dumps(error.value.report, allow_nan=False)


def test_saved_zero_kappa_is_valid_not_missing(tmp_path):
    caches = dict(G2=4, H2=1, I2=1, J2=2, K2=2, L2=0.5, M2=0.5, N2=0, O2="Rewrite")
    p = make_book(tmp_path / "zero.xlsx", [0, 0, 1, 1], [0, 1, 0, 1], caches)
    assert reliability(p)["status"] == "passed"


def test_single_constant_rater_can_have_defined_zero_kappa(tmp_path):
    p = make_book(tmp_path / "constant-rater.xlsx", [0, 0], [0, 1], caches=None)
    r = reliability(p, compute_only=True)
    assert r["codes"][0]["kappa"] == 0 and r["codes"][0]["status"] == "defined"


@pytest.mark.parametrize(
    "sheet,cell,value",
    [
        ("Reliability", "C2", None),
        ("Reliability", "D2", "1"),
        ("Reliability", "C2", True),
        ("Reliability", "C2", "=1"),
        ("Reliability", "C2", "#N/A"),
        ("Reliability", "B2", "UNKNOWN"),
        ("Reliability", "A2", "P01-999"),
        ("Reliability", "A2", "EX-01-001"),
        ("Reliability", "A3", "P01-001"),
        ("Reliability", "A602", "P01-001"),
        ("Reliability", "N2", "=1"),
        ("Reliability", "F2", "PEER"),
        ("Reliability", "C1", "rater_A"),
        ("Reliability", "N1", "agreement"),
        ("Passages", "A4", "P01-001"),
        ("Passages", "B3", "P02"),
        ("Cases", "D3", None),
        ("Cases", "A3", "P03"),
    ],
)
def test_malformed_workbooks_rejected_without_silently_dropping_rows(tmp_path, sheet, cell, value):
    p = make_book(tmp_path / "toy.xlsx")
    change(p, sheet, cell, value)
    with pytest.raises(ReliabilityInputError):
        reliability(p)


def test_same_passage_can_have_multiple_codes_without_pooling(tmp_path):
    p = make_book(tmp_path / "toy.xlsx")
    w = load_workbook(p)
    for row, values in enumerate([("P01-001", "PEER", 0, 1), ("P01-002", "PEER", 1, 0)], 20):
        for col, value in enumerate(values, 1):
            w["Reliability"].cell(row, col, value)
    w.save(p)
    w.close()
    r = reliability(p, compute_only=True)
    assert r["codes"][0]["kappa"] == 0.5 and r["codes"][1]["kappa"] == -1
    assert r["total_pairs"] == 10 and r["unique_passages"] == 8


def test_unreadable_input(tmp_path):
    p = tmp_path / "bad.xlsx"
    p.write_text("not a workbook")
    with pytest.raises(ReliabilityInputError):
        reliability(p)


def test_cli_exits_and_machine_readable_failures(tmp_path):
    p = make_book(tmp_path / "toy.xlsx")
    cmd = [sys.executable, "-m", "src.reliability", str(p), "--format", "json"]
    r = subprocess.run(cmd, capture_output=True, text=True)
    assert r.returncode == 0 and json.loads(r.stdout)["status"] == "passed"
    inject_test_caches(p, {"N2": 0.9})
    r = subprocess.run(cmd, capture_output=True, text=True)
    assert r.returncode == 1 and json.loads(r.stdout)["status"] == "failed"
    r = subprocess.run(cmd + ["--compute-only"], capture_output=True, text=True)
    assert r.returncode == 0 and json.loads(r.stdout)["status"] == "computed_only"
    change(p, "Reliability", "C2", None)
    r = subprocess.run(cmd, capture_output=True, text=True)
    assert r.returncode == 2 and "C2" in r.stderr
    r = subprocess.run(
        cmd[:3] + [str(TEMPLATE), "--format", "json"], capture_output=True, text=True
    )
    assert r.returncode == 1 and json.loads(r.stdout)["status"] == "no_data"
