"""M8 behavior tests use synthetic workbook copies, never research labels."""

from datetime import datetime
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import pymupdf
import pytest

from test_validate_sheet import case, passage, put, CODES, TEMPLATE
from src.pilot_report import (
    analyse_workbook,
    build_report,
    plot_rate_shock,
    ReportInputError,
    regime,
)


@pytest.fixture
def sample(tmp_path):
    w = load_workbook(TEMPLATE)
    for row, year in enumerate((2020, 2022, 2025), 3):
        cid = f"P{row - 2:02d}"
        case(w, row, cid, ("MA", "VT", "NH")[row - 3], year)
        put(
            w,
            "Cases",
            row,
            prior_authorized_roe=0.09,
            prior_order_date=datetime(year - 3, 1, 1),
            ust10_at_prior_order=0.015,
            ust10_at_order=0.04,
            perf_adjustment_bp=-47 if row == 3 else 0,
            advocate_rec=0.085,
            staff_rec=0.091,
            alj_rec=0.093,
            nat_avg_benchmark=0.094,
            minutes_spent=0 if row == 3 else 60,
        )
        passage(w, row, cid)
    passage(w, 6, "P01", passage_id="P01-002", method="Expected earnings", PEER=1)
    passage(
        w,
        7,
        "P02",
        passage_id="P02-002",
        section="Positions of parties",
        stance="Summarized only",
        PEER=1,
    )
    p = tmp_path / "synthetic.xlsx"
    w.save(p)
    w.close()
    return p


def mutate(p, sheet, row, **values):
    w = load_workbook(p)
    put(w, sheet, row, **values)
    w.save(p)
    w.close()


def analyse(p):
    return analyse_workbook(p, synthetic=True)


@pytest.mark.parametrize(
    "year,expected",
    [
        (2018, "Practice / out of frame"),
        (2019, "Pre (2019-21)"),
        (2021, "Pre (2019-21)"),
        (2022, "Transition (2022)"),
        (2023, "Post (2023-26)"),
        (2026, "Post (2023-26)"),
    ],
)
def test_written_date_regimes(year, expected):
    assert regime(datetime(year, 1, 1)) == expected


def test_numeric_units_and_source_preservation(sample):
    before = sample.read_bytes()
    r = analyse(sample)
    assert sample.read_bytes() == before
    assert r["n_cases"] == 3 and len(r["cases"]) == 3
    assert r["cases"][0]["awarded_roe_final"] == 0.0903
    assert r["scatter"]["n_plotted"] == 3
    assert r["scatter"]["points"][0]["delta_roe_pp"] == 0.5
    assert r["scatter"]["points"][0]["delta_ust10_pp"] == 2.5
    assert r["cases"][0]["sources_notes"].startswith("SYNTHETIC")


def test_missing_adjustment_is_not_zero(sample):
    mutate(sample, "Cases", 3, perf_adjustment_bp=None)
    c = analyse(sample)["cases"][0]
    assert c["awarded_roe"] == 0.095 and c["awarded_roe_final"] is None


@pytest.mark.parametrize(
    "field",
    [
        "prior_order_date",
        "prior_authorized_roe",
        "ust10_at_prior_order",
        "ust10_at_order",
        "awarded_roe",
        "order_date",
    ],
)
def test_scatter_excludes_missing_intervals_with_reasons(sample, field):
    mutate(sample, "Cases", 3, **{field: None})
    r = analyse(sample)
    assert r["scatter"]["n_plotted"] == 2 and r["scatter"]["n_excluded"] == 1
    assert field in r["scatter"]["excluded"][0]["missing_fields"]


def test_code_share_denominator_and_expected_earnings_union(sample):
    r = analyse(sample)
    assert r["coding"]["n_passages"] == 5 and r["coding"]["n_complete_commission"] == 4
    rows = {x["code"]: x for x in r["code_shares"] if x["group"] == "All"}
    assert (
        rows["PEER"]["count"] == 1
        and rows["PEER"]["denominator"] == 4
        and rows["PEER"]["share"] == 0.25
    )
    s = [x for x in r["sensitivity"] if x["group"] == "All"]
    assert s[0]["market_min"] == 4 and s[0]["benchmark_min"] == 1
    assert (
        s[1]["market_min"] == 3 and s[1]["benchmark_min"] == 1
    )  # EE and PEER same passage, count once


def test_unknown_mixed_method_produces_bounds(sample):
    mutate(sample, "Passages", 3, method="Multiple")
    s = [x for x in analyse(sample)["sensitivity"] if x["group"] == "All"][1]
    assert s["ambiguous_method"] == 1
    assert (s["market_min"], s["market_max"]) == (2, 3)
    assert (s["benchmark_min"], s["benchmark_max"]) == (1, 2)


def test_uncoded_commission_is_not_a_zero_code(sample):
    w = load_workbook(sample)
    for col, value in enumerate(
        ("P01-003", "P01", "1", "Commission analysis", "Synthetic uncoded text"), 1
    ):
        w["Passages"].cell(8, col, value)
    w.save(sample)
    w.close()
    r = analyse(sample)
    assert r["coding"]["n_commission"] == 5 and r["coding"]["n_complete_commission"] == 4
    assert r["coding"]["n_incomplete_commission"] == 1
    assert (
        next(x for x in r["code_shares"] if x["group"] == "All" and x["code"] == "MKT")["share"]
        == 1
    )


def test_minutes_zero_and_unknown_hit_rate(sample):
    mutate(sample, "Cases", 5, minutes_spent=None)
    r = analyse(sample)
    assert r["feasibility"]["median_minutes"] == 30 and r["feasibility"]["minutes_recorded"] == 2
    assert r["feasibility"]["m3_reviews"] == []
    assert all(x["rate"] is None for x in r["feasibility"]["m3_hit_rates"])


def add_reviews(p, rows):
    w = load_workbook(p)
    s = w.create_sheet("M3_Review")
    s.append(["case_id", "top_hit", "reviewer", "reviewer_type", "reviewed_on", "evidence"])
    for r in rows:
        s.append(r)
    w.save(p)
    w.close()


def test_hit_rates_are_separate_by_reviewer_type(sample):
    add_reviews(
        sample,
        [
            ["P01", 1, "Test assistant", "assistant", "2026-09-30", "Synthetic comparison"],
            ["P02", 0, "Test human", "human", "2026-09-30", "Synthetic comparison"],
        ],
    )
    rates = {x["reviewer_type"]: x for x in analyse(sample)["feasibility"]["m3_hit_rates"]}
    assert rates["assistant"]["rate"] == 1 and rates["human"]["rate"] == 0


@pytest.mark.parametrize(
    "row",
    [
        ["P01", 1, "", "assistant", "2026-09-30", "Evidence"],
        ["P01", True, "Test", "assistant", "2026-09-30", "Evidence"],
        ["P01", 1, "Test", "robot", "2026-09-30", "Evidence"],
        ["P15", 1, "Test", "assistant", "2026-09-30", "Evidence"],
        ["P01", 1, "Test", "assistant", "2026-99-99", "Evidence"],
        ["P01", 1, "Test", "assistant", "2026-09-30", None],
    ],
)
def test_malformed_hit_reviews_are_not_used(sample, row):
    add_reviews(sample, [row])
    with pytest.raises(ReportInputError):
        analyse(sample)


def test_invalid_m6_input_is_rejected(sample):
    mutate(sample, "Cases", 3, awarded_roe=9.5)
    with pytest.raises(ReportInputError, match="M6"):
        analyse(sample)


def test_m7_disagreement_stops_export(sample, tmp_path):
    from test_reliability import inject_test_caches

    inject_test_caches(sample, {"N2": 0.99})
    with pytest.raises(ReportInputError, match="disagree"):
        build_report(sample, tmp_path / "no-export", synthetic=True)
    assert not (tmp_path / "no-export").exists()


def test_synthetic_label_cannot_be_accidentally_omitted(sample):
    with pytest.raises(ReportInputError, match="synthetic"):
        analyse_workbook(sample)


def test_missing_numeric_sources_not_plotted(sample):
    mutate(sample, "Cases", 3, sources_notes=None)
    with pytest.raises(ReportInputError, match="sources"):
        analyse(sample)


def test_nontext_source_notes_do_not_crash_serialization(sample):
    mutate(sample, "Cases", 3, sources_notes=datetime(2020, 1, 1))
    with pytest.raises(ReportInputError, match="metadata"):
        analyse(sample)


def test_duplicate_hit_reviews_do_not_inflate_rate(sample):
    r = ["P01", 1, "Test", "assistant", "2026-09-30", "Synthetic comparison"]
    add_reviews(sample, [r, r])
    with pytest.raises(ReportInputError, match="duplicate"):
        analyse(sample)


def test_no_data_stays_missing_and_not_a_claim_of_15(sample):
    r = analyse_workbook(TEMPLATE)
    assert r["n_cases"] == 0 and r["scatter"]["n_plotted"] == 0
    assert r["coding"]["n_complete_commission"] == 0
    assert all(x["share"] is None for x in r["code_shares"])
    assert r["reliability"]["status"] == "no_data"


def test_cited_national_average_is_not_used_as_missing_benchmark(sample):
    mutate(sample, "Cases", 3, nat_avg_benchmark=None, nat_avg_value_cited=0.098)
    r = analyse(sample)
    assert r["cases"][0]["nat_avg_benchmark"] is None


def test_zero_roe_change_is_a_real_scatter_point(sample):
    mutate(sample, "Cases", 3, prior_authorized_roe=0.095, ust10_at_prior_order=0.04)
    point = analyse(sample)["scatter"]["points"][0]
    assert point["delta_roe_pp"] == 0 and point["delta_ust10_pp"] == 0


def test_scatter_uses_equal_units_and_actual_n(sample):
    from matplotlib import pyplot as plt

    r = analyse(sample)
    fig = plot_rate_shock(r)
    ax = fig.axes[0]
    assert "n = 3" in ax.get_title() and ax.get_aspect() == 1
    assert ax.get_xlim() == ax.get_ylim()
    assert list(ax.lines[0].get_xdata()) == list(ax.lines[0].get_ydata())
    plt.close(fig)


def test_one_command_bundle_and_pdf(sample, tmp_path):
    before = sample.read_bytes()
    out = tmp_path / "report"
    audit = build_report(sample, out, synthetic=True)
    assert sample.read_bytes() == before and audit["status"] == "complete"
    for name in [
        "report.html",
        "memo.md",
        "memo.pdf",
        "analysis.json",
        "tables/number_lines.csv",
        "tables/code_shares.csv",
        "tables/feasibility.csv",
        "tables/field_coverage.csv",
        "tables/reliability.csv",
        "tables/expected_earnings.csv",
        "figures/rate_shock.png",
        "figures/number_lines.svg",
        "figures/code_shares.png",
    ]:
        assert (out / name).is_file(), name
    with pymupdf.open(out / "memo.pdf") as pdf:
        assert len(pdf) == 2
        text = "".join(p.get_text() for p in pdf)
        assert "SYNTHETIC" in text and "n = 3" in text and "n = 15" not in text
    assert len(list((out / "figures/orders").glob("*.png"))) == 3
    html = (out / "report.html").read_text()
    assert "SYNTHETIC" in html and "Expected earnings" in html
    assert json.loads((out / "analysis.json").read_text())["n_cases"] == 3
    with pytest.raises(ReportInputError, match="exists"):
        build_report(sample, out, synthetic=True)


def test_export_escapes_markup_in_metadata(sample, tmp_path):
    mutate(sample, "Cases", 3, utility="<script>alert(1)</script>")
    out = tmp_path / "escaped"
    build_report(sample, out, synthetic=True)
    text = (out / "report.html").read_text()
    assert "<script>" not in text and "&lt;script&gt;" in text


def test_output_protected_or_symlink_rejected(sample, tmp_path):
    protected = tmp_path / "data/raw_orders/out"
    with pytest.raises(ReportInputError):
        build_report(sample, protected, synthetic=True)
    target = tmp_path / "real"
    target.mkdir()
    link = tmp_path / "link"
    link.symlink_to(target, target_is_directory=True)
    with pytest.raises(ReportInputError):
        build_report(sample, link / "out", synthetic=True)
    assert not list(target.iterdir())


def test_rebuild_is_deterministic(sample, tmp_path):
    a, b = tmp_path / "a", tmp_path / "b"
    build_report(sample, a, synthetic=True)
    build_report(sample, b, synthetic=True)
    for p in a.rglob("*"):
        if p.is_file():
            assert p.read_bytes() == (b / p.relative_to(a)).read_bytes(), p.name


def test_cli_invalid_input_does_not_create_output(tmp_path):
    out = tmp_path / "absent"
    p = subprocess.run(
        [sys.executable, "-m", "src.pilot_report", "missing.xlsx", str(out)],
        capture_output=True,
        text=True,
    )
    assert p.returncode == 2 and not out.exists()
