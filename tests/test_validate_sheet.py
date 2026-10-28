"""M6 checks use disposable copies of the supplied workbook, never live inputs."""

from datetime import datetime
from pathlib import Path
import json
import subprocess
import sys

from openpyxl import load_workbook
import pytest

from src.validate_sheet import ValidationInputError, validate_sheet

TEMPLATE = (
    Path(__file__).resolve().parents[1]
    / "project_docs/Anchored_Returns_Pilot_Coding_Sheet_v0.1.2.xlsx"
)
CODES = "MKT PEER OWN GRAD LAW CAP RISK AFF PERF DEF CONC".split()


def put(book, sheet, row, **values):
    columns = {}
    for cell in book[sheet][1]:
        if cell.value:
            columns.setdefault(cell.value, cell.column)
    for key, value in values.items():
        book[sheet].cell(row, columns[key]).value = value


def case(book, row=3, case_id="P01", state="NH", year=2025):
    put(
        book,
        "Cases",
        row,
        case_id=case_id,
        sample_status="Included by design" if state == "NH" else "Sample",
        state=state,
        commission="Synthetic commission",
        docket="DE 24-070" if state == "NH" else "TEST",
        utility="Synthetic utility",
        service="Electric",
        segment="Distribution",
        decision_type="Fully litigated",
        order_date=datetime(year, 7, 25),
        order_url="https://example.org/synthetic.pdf",
        roe_section_pages="1-3",
        requested_roe_initial=0.105,
        requested_roe=0.10,
        awarded_roe=0.095,
        decisive_rationale="MKT",
        conclusory_flag="N",
        coder="SYNTHETIC",
        date_coded=datetime(2026, 9, 30),
        sources_notes="SYNTHETIC TEST ONLY; PDF p. 2.",
    )


def passage(book, row=3, case_id="P01", **changes):
    values = dict(
        passage_id=case_id + "-001",
        case_id=case_id,
        page_para="PDF p. 2",
        excerpt="Synthetic text for software tests only.",
        section="Commission analysis",
        method="DCF",
        primary_code="MKT",
        stance="Relied on",
        coder="SYNTHETIC",
        **{code: int(code == "MKT") for code in CODES},
    )
    values.update(changes)
    put(book, "Passages", row, **values)


@pytest.fixture
def book():
    w = load_workbook(TEMPLATE)
    case(w)
    passage(w)
    yield w
    w.close()


def check(tmp_path, book, **options):
    path = tmp_path / "check.xlsx"
    book.save(path)
    before = path.read_bytes()
    result = validate_sheet(path, **options)
    assert path.read_bytes() == before
    return result


def errors(result):
    return {i["code"] for i in result["issues"] if i["severity"] == "error"}


def test_valid_populated_case_preserves_source(tmp_path, book):
    r = check(tmp_path, book)
    assert not errors(r)
    assert r["counts"]["cases"] == 1 and r["counts"]["coded_cases"] == 1
    assert r["counts"]["example_cases_excluded"] == 1
    assert r["counts"]["placeholder_cases_excluded"] == 14


@pytest.mark.parametrize(
    "sheet,field,value,code",
    [
        ("Cases", "state", "NY", "invalid_choice"),
        ("Cases", "sample_status", "Excluded", "invalid_choice"),
        ("Cases", "service", "Water", "invalid_choice"),
        ("Passages", "section", "Commission", "invalid_choice"),
        ("Passages", "method", "CAPM", "invalid_choice"),
        ("Passages", "stance", "Accepted", "invalid_choice"),
        ("Passages", "PEER", 2, "invalid_binary"),
        ("Passages", "MKT", "1", "invalid_binary"),
        ("Passages", "MKT", True, "invalid_binary"),
        ("Cases", "awarded_roe", 9.5, "roe_range"),
        ("Cases", "awarded_roe", 0, "roe_range"),
        ("Cases", "staff_rec", "9.5%", "invalid_number"),
        ("Cases", "awarded_roe", True, "invalid_number"),
        ("Cases", "requested_roe", "=1/10", "input_formula"),
        ("Cases", "awarded_roe", "#DIV/0!", "input_error"),
    ],
)
def test_bad_values_are_not_coerced(tmp_path, book, sheet, field, value, code):
    put(book, sheet, 3, **{field: value})
    assert code in errors(check(tmp_path, book))


@pytest.mark.parametrize(
    "field",
    [
        "requested_roe_initial",
        "requested_roe",
        "utility_witness_low",
        "utility_witness_high",
        "staff_rec",
        "advocate_rec",
        "other_intervenor_rec",
        "alj_rec",
        "commission_range_low",
        "commission_range_high",
        "awarded_roe",
        "prior_authorized_roe",
        "nat_avg_value_cited",
        "nat_avg_benchmark",
    ],
)
def test_all_roe_input_fields_guard_percentage_units(tmp_path, book, field):
    put(book, "Cases", 3, **{field: 9.5})
    assert any(
        i["code"] == "roe_range" and i["field"] == field for i in check(tmp_path, book)["issues"]
    )


def test_roe_endpoints_and_missing_values_are_valid(tmp_path, book):
    put(
        book,
        "Cases",
        3,
        utility_witness_low=0.03,
        utility_witness_high=0.20,
        staff_rec=None,
        advocate_rec=None,
    )
    assert not errors(check(tmp_path, book))


def test_primary_flag_must_equal_one(tmp_path, book):
    put(book, "Passages", 3, MKT=0)
    assert "primary_code_unset" in errors(check(tmp_path, book))


def test_multiple_codes_can_legitimately_cooccur(tmp_path, book):
    put(book, "Passages", 3, PEER=1, OWN=1)
    assert not errors(check(tmp_path, book))


def test_every_coded_case_needs_coded_commission_analysis(tmp_path, book):
    put(book, "Passages", 3, section="Positions of parties", stance="Summarized only")
    assert "no_commission_analysis" in errors(check(tmp_path, book))
    passage(book, 4, passage_id="P01-002")
    assert not errors(check(tmp_path, book))


def test_uncoded_commission_heading_does_not_satisfy_requirement(tmp_path, book):
    for key in CODES + ["primary_code", "stance", "coder", "method"]:
        put(book, "Passages", 3, **{key: None})
    assert "no_commission_analysis" in errors(check(tmp_path, book))


def test_summarized_only_is_forbidden_in_commission_analysis(tmp_path, book):
    put(book, "Passages", 3, stance="Summarized only")
    assert "summarized_commission" in errors(check(tmp_path, book))


def test_duplicate_and_foreign_ids(tmp_path, book):
    passage(book, 4)
    assert "duplicate_passage_id" in errors(check(tmp_path, book))
    put(book, "Passages", 4, passage_id="P010-001", case_id="P01")
    assert "passage_prefix" in errors(check(tmp_path, book))
    put(book, "Passages", 4, case_id="P02", passage_id="P02-001")
    assert "unknown_case_id" in errors(check(tmp_path, book))


def test_duplicate_cases_and_orphan_cells(tmp_path, book):
    case(book, 4)
    assert "duplicate_case_id" in errors(check(tmp_path, book))
    put(book, "Cases", 4, case_id=None)
    assert "missing_case_id" in errors(check(tmp_path, book))


def test_request_increase_needs_specific_explanation(tmp_path, book):
    put(book, "Cases", 3, requested_roe_initial=0.09, sources_notes="Some unrelated notes.")
    assert "request_increase_unexplained" in errors(check(tmp_path, book))
    put(book, "Cases", 3, sources_notes="request_revision_explanation:")
    assert "request_increase_unexplained" in errors(check(tmp_path, book))
    put(
        book,
        "Cases",
        3,
        sources_notes="PDF p. 2. request_revision_explanation: Company revised its request upward in rebuttal.",
    )
    r = check(tmp_path, book)
    assert "request_increase_unexplained" not in errors(r)
    assert any(i["code"] == "request_increase_noted" for i in r["issues"])


def test_draft_blanks_are_warnings_complete_blanks_are_errors(tmp_path):
    w = load_workbook(TEMPLATE)
    put(w, "Cases", 3, state="NH", order_date=datetime(2025, 7, 25))
    put(
        w,
        "Passages",
        3,
        passage_id="P01-001",
        case_id="P01",
        page_para="PDF p. 2",
        excerpt="Uncoded draft",
    )
    r = check(tmp_path, w)
    assert not errors(r) and r["status"] == "incomplete"
    assert "incomplete_passage" in errors(check(tmp_path, w, mode="complete"))
    assert "no_commission_analysis" not in errors(r)


def test_empty_template_is_not_a_complete_pilot(tmp_path):
    w = load_workbook(TEMPLATE)
    r = check(tmp_path, w)
    assert r["counts"]["cases"] == 0 and r["status"] == "incomplete"
    assert r["sample"]["pre"] == r["sample"]["post"] == 0
    assert "empty_workbook" in errors(check(tmp_path, w, mode="complete"))


def test_modified_example_cannot_escape_checks(tmp_path):
    w = load_workbook(TEMPLATE)
    w["Cases"]["X2"] = 9.5
    assert "example_modified" in errors(check(tmp_path, w))
    w = load_workbook(TEMPLATE)
    w["Passages"]["E2"] = "Real observation mislabeled as example"
    assert "example_modified" in errors(check(tmp_path, w))


def test_date_boundaries_and_quota_counts_ignore_cache(tmp_path, book):
    years = [2019] * 5 + [2022] * 5 + [2023] * 5
    states = ["CT", "MA", "ME", "NH", "RI", "VT"] * 3
    for i, year in enumerate(years):
        case(book, i + 3, f"P{i + 1:02}", states[i], year)
        passage(book, i + 3, f"P{i + 1:02}")
    r = check(tmp_path, book)
    assert r["sample"]["pre"] == 5 and r["sample"]["transition"] == 5 and r["sample"]["post"] == 5
    assert not any(i["code"] in ("quota_pre", "quota_post") for i in r["issues"])
    put(book, "Cases", 3, order_date=datetime(2022, 1, 1))
    assert any(
        i["code"] == "quota_pre" and i["severity"] == "warning"
        for i in check(tmp_path, book)["issues"]
    )


def test_quotas_wait_until_fifteen_entered(tmp_path, book):
    for i in range(14):
        case(book, i + 3, f"P{i + 1:02}", year=2025)
        passage(book, i + 3, f"P{i + 1:02}")
    assert not any(i["code"] == "quota_pre" for i in check(tmp_path, book)["issues"])
    case(book, 17, "P15")
    passage(book, 17, "P15")
    r = check(tmp_path, book)
    assert {"quota_pre", "quota_vermont", "quota_states"} <= {i["code"] for i in r["issues"]}


@pytest.mark.parametrize(
    "value", [datetime(2018, 12, 31), datetime(2027, 1, 1), "2025-02-30", 44800]
)
def test_invalid_or_out_of_frame_dates(tmp_path, book, value):
    put(book, "Cases", 3, order_date=value)
    if isinstance(value, int):
        book["Cases"]["K3"].number_format = "General"
    assert errors(check(tmp_path, book)) & {"invalid_date", "sample_date"}


def test_practice_scope_does_not_apply_pilot_quotas(tmp_path, book):
    case(book, case_id="PRACTICE-MA-17-05", state="MA", year=2017)
    passage(book, case_id="PRACTICE-MA-17-05")
    r = check(tmp_path, book, scope="practice")
    assert not errors(r)
    assert r["sample"]["quotas_applied"] is False


def test_ranges_and_performance_adjustment_units(tmp_path, book):
    put(book, "Cases", 3, awarded_roe=0.091, perf_adjustment_bp=-47)
    r = check(tmp_path, book)
    assert r["derived_checks"][0]["awarded_roe_final"] == "0.0863"
    put(
        book,
        "Cases",
        3,
        perf_adjustment_bp=-1000,
        utility_witness_low=0.11,
        utility_witness_high=0.10,
    )
    assert {"roe_range", "inverted_range"} <= errors(check(tmp_path, book))


def test_input_formula_and_overwritten_derived_formula(tmp_path, book):
    book["Cases"]["AW3"] = "=X3+Y3/100"
    assert "derived_formula_changed" in errors(check(tmp_path, book))


def test_dropdown_changes_cannot_whitelist_bad_values(tmp_path, book):
    book["Lists"]["A2"] = "NY"
    put(book, "Cases", 3, state="NY")
    assert {"list_schema_changed", "invalid_choice"} <= errors(check(tmp_path, book))
    book["Passages"].data_validations.dataValidation = []
    put(book, "Passages", 3, PEER=7)
    assert "invalid_binary" in errors(check(tmp_path, book))


def test_rows_outside_formula_capacity_are_not_silently_skipped(tmp_path, book):
    passage(book, 401, passage_id="P01-002", PEER=2)
    assert {"outside_template_capacity", "invalid_binary"} <= errors(check(tmp_path, book))


def test_reliability_inputs_are_validated_but_kappa_not_certified(tmp_path, book):
    put(book, "Reliability", 2, passage_id="P01-001", code="BOGUS", coding_A=2, coding_B=1)
    assert {"invalid_choice", "invalid_binary"} <= errors(check(tmp_path, book))
    put(book, "Reliability", 2, code="MKT", coding_A=0)
    put(book, "Reliability", 3, passage_id="P01-001", code="MKT", coding_A=1, coding_B=0)
    assert "duplicate_reliability_pair" in errors(check(tmp_path, book))


def test_missing_headers_and_bad_file(tmp_path, book):
    book["Cases"]["X1"] = "awarded_RETURN"
    assert "header_schema" in errors(check(tmp_path, book))
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a spreadsheet")
    with pytest.raises(ValidationInputError):
        validate_sheet(bad)


def test_research_case_cannot_occupy_excluded_example_row(tmp_path, book):
    case(book, 2, case_id="P02")
    assert "reserved_example_row" in errors(check(tmp_path, book))


def test_unknown_extra_input_column_is_not_ignored(tmp_path, book):
    book["Cases"]["BF1"] = "another_roe"
    book["Cases"]["BF3"] = 9.5
    assert "header_schema" in errors(check(tmp_path, book))


def test_invalid_id_keeps_json_diagnostics_serializable(tmp_path, book):
    put(book, "Cases", 3, case_id=datetime(2025, 1, 1))
    r = check(tmp_path, book)
    assert "invalid_case_id" in errors(r)
    json.dumps(r, allow_nan=False)


def test_cli_json_exit_status_and_no_writes(tmp_path, book):
    path = tmp_path / "valid.xlsx"
    book.save(path)
    before = path.read_bytes()
    cmd = [sys.executable, "-m", "src.validate_sheet", str(path), "--format", "json"]
    r = subprocess.run(cmd, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert json.loads(r.stdout)["counts"]["cases"] == 1
    strict = subprocess.run(cmd + ["--fail-on-warnings"], capture_output=True, text=True)
    assert strict.returncode == 1 and json.loads(strict.stdout)["errors"] == 0
    put(book, "Cases", 3, awarded_roe=9.5)
    book.save(path)
    before = path.read_bytes()
    r = subprocess.run(cmd, capture_output=True, text=True)
    assert r.returncode == 1 and "roe_range" in r.stdout and path.read_bytes() == before
    r = subprocess.run(cmd + ["--mode", "unknown"], capture_output=True, text=True)
    assert r.returncode == 2
