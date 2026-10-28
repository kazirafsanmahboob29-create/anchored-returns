"""M6: read-only structural and consistency checks for the v0.1.2 coding sheet.

This does not evaluate Excel formulas, verify source facts, assign codes, or certify
research readiness. Draft mode reports unfinished work; complete mode rejects it.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import re
import sys
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

SCHEMA_PATH = Path(__file__).with_name("workbook_schema_v012.json")
SCHEMA = json.loads(SCHEMA_PATH.read_text())
VERSION = "1.0"
CODES = tuple(SCHEMA["lists"]["Code"])
ROE_FIELDS = (
    "requested_roe_initial",
    "requested_roe",
    "utility_witness_low",
    "utility_witness_high",
    "staff_rec",
    "advocate_rec",
    "other_intervenor_rec",
    "alj_rec",
    "commission_range_low",
    "commission_range_high",
    "awarded_roe",
    "prior_authorized_roe",
    "nat_avg_value_cited",
    "nat_avg_benchmark",
)
DATE_FIELDS = ("open_meeting_date", "order_date", "prior_order_date", "date_coded")
CASE_CHOICES = dict(
    sample_status="SampleStatus",
    state="State",
    service="Service",
    segment="Segment",
    decision_type="Decision",
    nat_avg_cited_in_order="YN",
    statutory_peer_mandate="YN",
    decisive_rationale="Code",
    conclusory_flag="YN",
    commissioner_selection="Selection",
)
PASSAGE_CHOICES = dict(section="Section", method="Method", primary_code="Code", stance="Stance")
CASE_METADATA = (
    "sample_status",
    "state",
    "commission",
    "docket",
    "utility",
    "service",
    "segment",
    "decision_type",
    "order_date",
    "order_url",
    "roe_section_pages",
)
CASE_CODING = ("decisive_rationale", "conclusory_flag", "coder", "date_coded")
ID = re.compile(r"[A-Za-z0-9][A-Za-z0-9_.-]*\Z")
PILOT_ID = re.compile(r"P(?:0[1-9]|1[0-5])\Z")


class ValidationInputError(ValueError):
    """Unreadable input or unsupported invocation; never return an empty pass."""


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value):
    # bool is an int subclass, but an Excel logical is not a research number.
    if type(value) in (int, float) and math.isfinite(value):
        return Decimal(str(value))
    return None


def _date(value):
    if type(value) is datetime:
        return value.date() if value.time().isoformat() == "00:00:00" else None
    if type(value) is date:
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        try:
            return date.fromisoformat(value)
        except ValueError:
            pass
    return None


def _signature(cells):
    values = [
        c.value.isoformat() if isinstance(c.value, (date, datetime)) else c.value for c in cells
    ]
    return hashlib.sha256(
        json.dumps(values, ensure_ascii=False, separators=(",", ":"), default=str).encode()
    ).hexdigest()


class _Checks:
    def __init__(self, workbook, mode, scope):
        self.book, self.mode, self.scope = workbook, mode, scope
        self.issues, self.derived = [], []
        self.counts = dict(
            cases=0,
            passages=0,
            coded_cases=0,
            coded_passages=0,
            example_cases_excluded=0,
            example_passages_excluded=0,
            placeholder_cases_excluded=0,
            reliability_rows=0,
        )
        self.sample = dict(quotas_applied=False, pre=0, transition=0, post=0, states={})
        self.maps = {}
        self.cases, self.passages, self.coded_cases, self.commission_cases = {}, set(), set(), set()

    def issue(self, code, message, sheet=None, row=None, field=None, severity="error"):
        col = self.maps.get(sheet, {}).get(field)
        cell = f"{get_column_letter(col)}{row}" if col and row else None
        self.issues.append(
            dict(
                severity=severity,
                code=code,
                sheet=sheet,
                row=row,
                field=field,
                cell=cell,
                message=message,
            )
        )

    def missing(self, code, fields, sheet, row=None):
        if fields:
            self.issue(
                code,
                "Unfinished fields: " + ", ".join(fields),
                sheet,
                row,
                severity="error" if self.mode == "complete" else "incomplete",
            )

    def schema(self):
        for sheet, key in [
            ("Cases", "case_headers"),
            ("Passages", "passage_headers"),
            ("Reliability", "reliability_headers"),
        ]:
            headers = SCHEMA[key]
            if sheet not in self.book:
                self.issue("missing_sheet", f"Missing required worksheet {sheet}.", sheet)
                continue
            ws = self.book[sheet]
            if ws.max_row > 10000 or ws.max_column > 256:
                raise ValidationInputError(
                    f"{sheet}: exceeds supported 10,000-row/256-column limit."
                )
            if [ws.cell(1, i).value for i in range(1, len(headers) + 1)] != headers:
                self.issue(
                    "header_schema",
                    f"Expected v0.1.2 {sheet} headers in their original positions.",
                    sheet,
                    1,
                )
            if sheet != "Reliability" and any(
                not _blank(ws.cell(1, i).value) for i in range(len(headers) + 1, ws.max_column + 1)
            ):
                self.issue(
                    "header_schema",
                    "Unexpected input columns; update the versioned schema before adding fields.",
                    sheet,
                    1,
                )
            self.maps[sheet] = {v: i for i, v in enumerate(headers, 1)}
        if "Lists" not in self.book:
            self.issue("missing_sheet", "Missing Lists worksheet.", "Lists")
        else:
            ws = self.book["Lists"]
            if ws.max_row > 10000 or ws.max_column > 256:
                raise ValidationInputError("Lists exceeds supported dimensions.")
            actual = {
                ws.cell(1, i).value: [
                    ws.cell(r, i).value
                    for r in range(2, ws.max_row + 1)
                    if not _blank(ws.cell(r, i).value)
                ]
                for i in range(1, ws.max_column + 1)
                if not _blank(ws.cell(1, i).value)
            }
            if actual != SCHEMA["lists"]:
                self.issue(
                    "list_schema_changed",
                    "Lists differs from codebook v0.1.2; version the schema before changing categories.",
                    "Lists",
                )
        return not any(i["code"] in ("missing_sheet", "header_schema") for i in self.issues)

    def values(self, sheet, row):
        return {key: self.book[sheet].cell(row, col) for key, col in self.maps[sheet].items()}

    def typed_inputs(self, sheet, row, cells, choices, binary=()):
        for key, cell in cells.items():
            value = cell.value
            if _blank(value):
                continue
            if cell.data_type == "f" or cell.data_type == "e":
                kind = "formula" if cell.data_type == "f" else "error"
                self.issue(
                    "input_" + kind,
                    f"Input contains an Excel {kind}; enter the source value explicitly.",
                    sheet,
                    row,
                    key,
                )
                continue
            if key in binary and (_number(value) is None or value not in (0, 1)):
                self.issue(
                    "invalid_binary",
                    "Code must be numeric 0 or 1; text and logical values are invalid.",
                    sheet,
                    row,
                    key,
                )
            elif key in choices and value not in SCHEMA["lists"][choices[key]]:
                self.issue(
                    "invalid_choice",
                    f"Allowed values: {', '.join(map(str, SCHEMA['lists'][choices[key]]))}.",
                    sheet,
                    row,
                    key,
                )
        # The independent checks above work even where dropdowns were removed.
        populated = [
            c for k, c in cells.items() if k in (*choices, *binary) and not _blank(c.value)
        ]
        missing = [
            c.coordinate
            for c in populated
            if not any(
                d.type == "list" and c.coordinate in d.sqref
                for d in self.book[sheet].data_validations.dataValidation
            )
        ]
        if missing:
            self.issue(
                "dropdown_coverage_missing",
                "No list dropdown at "
                + ", ".join(missing)
                + "; values still checked independently.",
                sheet,
                row,
                severity="warning",
            )

    def numeric(self, sheet, row, cells, fields, minimum=None, maximum=None, code="invalid_number"):
        for field in fields:
            cell = cells[field]
            if _blank(cell.value) or cell.data_type in ("f", "e"):
                continue
            number = _number(cell.value)
            if number is None:
                self.issue(
                    "invalid_number",
                    "Expected a finite numeric value, not text or a logical value.",
                    sheet,
                    row,
                    field,
                )
            elif (minimum is not None and number < minimum) or (
                maximum is not None and number > maximum
            ):
                self.issue(
                    code,
                    f"Expected {minimum} through {maximum}; ROEs/rates use fractions, not whole percentages.",
                    sheet,
                    row,
                    field,
                )

    def case_rows(self):
        seen = set()
        for row in range(2, self.book["Cases"].max_row + 1):
            all_cells = self.values("Cases", row)
            cells = dict(list(all_cells.items())[:41])
            value = {k: c.value for k, c in cells.items()}
            if all(_blank(v) for v in value.values()):
                continue
            cid = value["case_id"]
            if isinstance(cid, str):
                if cid in seen:
                    self.issue(
                        "duplicate_case_id", f"Duplicate case ID {cid}.", "Cases", row, "case_id"
                    )
                seen.add(cid)
            if cid == "EX-01":
                if _signature(list(cells.values())) != SCHEMA["example_signatures"]["Cases"]:
                    self.issue(
                        "example_modified",
                        "Reserved illustrative row changed; do not disguise research data as EX-01.",
                        "Cases",
                        row,
                    )
                else:
                    self.counts["example_cases_excluded"] += 1
                continue
            if (
                isinstance(cid, str)
                and PILOT_ID.fullmatch(cid)
                and all(_blank(v) for k, v in value.items() if k != "case_id")
            ):
                self.counts["placeholder_cases_excluded"] += 1
                continue
            self.counts["cases"] += 1
            if row == 2:
                self.issue(
                    "reserved_example_row",
                    "Cases row 2 is reserved for the example and excluded by Summary. Put research/practice cases in rows 3–17.",
                    "Cases",
                    row,
                )
            if row > 17:
                self.issue(
                    "outside_template_capacity",
                    "Cases row exceeds the template summary/formula capacity (row 17).",
                    "Cases",
                    row,
                )
            if not isinstance(cid, str) or not ID.fullmatch(cid):
                self.issue(
                    "missing_case_id" if _blank(cid) else "invalid_case_id",
                    "Supply a unique case ID with letters, digits, dots, underscores or hyphens.",
                    "Cases",
                    row,
                    "case_id",
                )
            else:
                self.cases.setdefault(cid, (row, value))
                if self.scope == "pilot" and not PILOT_ID.fullmatch(cid):
                    self.issue(
                        "pilot_case_id",
                        "Pilot workbook uses P01–P15. Use --scope practice for separate practice records.",
                        "Cases",
                        row,
                        "case_id",
                    )
            self.typed_inputs("Cases", row, cells, CASE_CHOICES)
            self.numeric(
                "Cases", row, cells, ROE_FIELDS, Decimal(".03"), Decimal(".20"), "roe_range"
            )
            self.numeric("Cases", row, cells, ("perf_adjustment_bp",))
            self.numeric("Cases", row, cells, ("equity_ratio_awarded",), Decimal("0"), Decimal("1"))
            self.numeric(
                "Cases",
                row,
                cells,
                ("ust10_at_order", "ust10_at_prior_order"),
                Decimal("-1"),
                Decimal("1"),
            )
            self.numeric("Cases", row, cells, ("minutes_spent",), Decimal("0"))
            for field in DATE_FIELDS:
                if not _blank(value[field]) and _date(value[field]) is None:
                    self.issue(
                        "invalid_date",
                        "Use a real Excel date or exact YYYY-MM-DD date; no raw serial or timestamp.",
                        "Cases",
                        row,
                        field,
                    )
            written = _date(value["order_date"])
            if written and self.scope == "pilot" and not 2019 <= written.year <= 2026:
                self.issue(
                    "sample_date",
                    "Written order date is outside 2019–2026; it cannot count toward period quotas.",
                    "Cases",
                    row,
                    "order_date",
                )
            if written and self.scope == "practice" and 2019 <= written.year <= 2026:
                self.issue(
                    "practice_in_sample_window",
                    "Practice should use an out-of-frame order.",
                    "Cases",
                    row,
                    "order_date",
                    "warning",
                )
            prior = _date(value["prior_order_date"])
            if prior and written and prior >= written:
                self.issue(
                    "prior_date_order",
                    "Prior order must precede this written order.",
                    "Cases",
                    row,
                    "prior_order_date",
                )
            required_metadata = [
                f for f in CASE_METADATA if self.scope == "pilot" or f != "sample_status"
            ]
            self.missing(
                "incomplete_case_metadata",
                [f for f in required_metadata if _blank(value[f])],
                "Cases",
                row,
            )
            coded = any(not _blank(value[f]) for f in CASE_CODING)
            if coded or self.mode == "complete":
                if isinstance(cid, str):
                    self.coded_cases.add(cid)
                self.missing(
                    "incomplete_case_coding",
                    [f for f in CASE_CODING if _blank(value[f])],
                    "Cases",
                    row,
                )
            if any(
                not _blank(value[f])
                for f in (
                    *ROE_FIELDS,
                    "perf_adjustment_bp",
                    "ust10_at_order",
                    "ust10_at_prior_order",
                    "equity_ratio_awarded",
                )
            ) and _blank(value["sources_notes"]):
                self.missing(
                    "missing_numeric_sources",
                    ["sources_notes (identify pages for entered numbers)"],
                    "Cases",
                    row,
                )
            initial, final = (
                _number(value["requested_roe_initial"]),
                _number(value["requested_roe"]),
            )
            if initial is not None and final is not None and initial < final:
                note = str(value["sources_notes"] or "")
                explained = re.search(r"(?im)\brequest_revision_explanation:\s*(\S[^\n]*)", note)
                self.issue(
                    "request_increase_noted" if explained else "request_increase_unexplained",
                    "Initial request is below final. "
                    + (
                        "Explanation recorded; its substantive adequacy is not machine-verified."
                        if explained
                        else "Add request_revision_explanation: followed by the reason and source in sources_notes."
                    ),
                    "Cases",
                    row,
                    "requested_roe_initial",
                    "warning" if explained else "error",
                )
            for low, high in [
                ("utility_witness_low", "utility_witness_high"),
                ("commission_range_low", "commission_range_high"),
            ]:
                a, b = _number(value[low]), _number(value[high])
                if a is not None and b is not None and a > b:
                    self.issue("inverted_range", f"{low} exceeds {high}.", "Cases", row, low)
            base, bp = _number(value["awarded_roe"]), _number(value["perf_adjustment_bp"])
            if base is not None and (_blank(value["perf_adjustment_bp"]) or bp is not None):
                adjusted = base + (bp if bp is not None else Decimal(0)) / 10000
                self.derived.append(
                    dict(
                        case_id=str(cid) if cid is not None else None,
                        row=row,
                        awarded_roe_final=str(adjusted),
                        blank_adjustment_treated_as_zero_for_template_check=bp is None,
                    )
                )
                if not Decimal(".03") <= adjusted <= Decimal(".20"):
                    self.issue(
                        "roe_range",
                        "Base plus signed basis-point adjustment is outside 0.03–0.20.",
                        "Cases",
                        row,
                        "awarded_roe_final",
                    )
            for origin, formula in SCHEMA["case_formulas"].items():
                col = re.match(r"[A-Z]+", origin)[0]
                target = f"{col}{row}"
                expected = Translator(formula, origin=origin).translate_formula(target)
                if self.book["Cases"][target].value != expected:
                    field = SCHEMA["case_headers"][self.book["Cases"][target].column - 1]
                    self.issue(
                        "derived_formula_changed",
                        "Expected the v0.1.2 row-relative formula; formula results are not trusted as inputs.",
                        "Cases",
                        row,
                        field,
                    )

    def passage_rows(self):
        seen = set()
        for row in range(2, self.book["Passages"].max_row + 1):
            cells = self.values("Passages", row)
            value = {k: c.value for k, c in cells.items()}
            if all(_blank(v) for v in value.values()):
                continue
            pid, cid = value["passage_id"], value["case_id"]
            if isinstance(pid, str):
                if pid in seen:
                    self.issue(
                        "duplicate_passage_id",
                        f"Duplicate passage ID {pid}.",
                        "Passages",
                        row,
                        "passage_id",
                    )
                seen.add(pid)
            if pid == "EX-01-001" or cid == "EX-01":
                if _signature(list(cells.values())) != SCHEMA["example_signatures"]["Passages"]:
                    self.issue(
                        "example_modified",
                        "Reserved example passage changed; do not exclude real observations as examples.",
                        "Passages",
                        row,
                    )
                else:
                    self.counts["example_passages_excluded"] += 1
                continue
            self.counts["passages"] += 1
            if row > 400:
                self.issue(
                    "outside_template_capacity",
                    "Passage is beyond row 400 and would be omitted by template summaries.",
                    "Passages",
                    row,
                )
            self.typed_inputs("Passages", row, cells, PASSAGE_CHOICES, CODES)
            if not isinstance(pid, str) or not ID.fullmatch(pid):
                self.issue(
                    "invalid_passage_id",
                    "Supply a nonblank, well-formed passage ID.",
                    "Passages",
                    row,
                    "passage_id",
                )
            elif isinstance(cid, str) and (
                not pid.startswith(cid + "-") or len(pid) == len(cid) + 1
            ):
                self.issue(
                    "passage_prefix",
                    "Passage ID must begin with its exact case ID followed by a hyphen and suffix.",
                    "Passages",
                    row,
                    "passage_id",
                )
            if not isinstance(cid, str) or cid not in self.cases:
                self.issue(
                    "unknown_case_id",
                    "Passage must reference a populated Cases row; empty P01–P15 placeholders are not case records.",
                    "Passages",
                    row,
                    "case_id",
                )
            if isinstance(pid, str):
                self.passages.add(pid)
            self.missing(
                "incomplete_passage_text",
                [f for f in ("excerpt", "page_para") if _blank(value[f])],
                "Passages",
                row,
            )
            coding_fields = (*CODES, "primary_code", "stance", "method", "coder")
            coded = any(not _blank(value[f]) for f in coding_fields)
            if coded and isinstance(cid, str):
                self.coded_cases.add(cid)
                self.counts["coded_passages"] += 1
            required = ("section", *CODES, "primary_code", "stance", "coder")
            missing = [f for f in required if _blank(value[f])]
            if value["MKT"] == 1 and _blank(value["method"]):
                missing.append("method")
            # Report uncoded rows too, but draft mode never converts blanks to zeroes.
            self.missing("incomplete_passage", missing, "Passages", row)
            primary = value["primary_code"]
            if primary in CODES and (_number(value[primary]) != Decimal(1)):
                self.issue(
                    "primary_code_unset",
                    f"Primary code {primary} must have its own flag equal to 1.",
                    "Passages",
                    row,
                    "primary_code",
                )
            if value["section"] == "Commission analysis" and value["stance"] == "Summarized only":
                self.issue(
                    "summarized_commission",
                    "Summarized only is not allowed in Commission analysis.",
                    "Passages",
                    row,
                    "stance",
                )
            if (
                coded
                and not missing
                and value["section"] == "Commission analysis"
                and primary in CODES
                and _number(value[primary]) == 1
                and value["stance"] in ("Relied on", "Considered", "Rejected")
                and all(_number(value[c]) in (Decimal(0), Decimal(1)) for c in CODES)
            ):
                self.commission_cases.add(cid)

    def reliability_rows(self):
        seen = set()
        for row in range(2, self.book["Reliability"].max_row + 1):
            cells = self.values("Reliability", row)
            value = {k: c.value for k, c in cells.items()}
            if all(_blank(v) for v in value.values()):
                continue
            self.counts["reliability_rows"] += 1
            self.typed_inputs("Reliability", row, cells, {"code": "Code"}, ("coding_A", "coding_B"))
            if row > 601:
                self.issue(
                    "outside_template_capacity",
                    "Reliability row exceeds formula capacity (row 601).",
                    "Reliability",
                    row,
                )
            pid, code = value["passage_id"], value["code"]
            if not isinstance(pid, str) or pid not in self.passages:
                self.issue(
                    "unknown_passage_id",
                    "Reliability row must reference an existing non-example passage.",
                    "Reliability",
                    row,
                    "passage_id",
                )
            if isinstance(pid, str) and isinstance(code, str):
                if (pid, code) in seen:
                    self.issue(
                        "duplicate_reliability_pair",
                        "Only one A/B pair per passage and code is supported.",
                        "Reliability",
                        row,
                    )
                seen.add((pid, code))
            self.missing(
                "incomplete_reliability_pair",
                [f for f in value if _blank(value[f])],
                "Reliability",
                row,
            )

    def cross_checks(self):
        self.counts["coded_cases"] = len(self.coded_cases & self.cases.keys())
        for cid in sorted(self.coded_cases & self.cases.keys()):
            if cid not in self.commission_cases:
                self.issue(
                    "no_commission_analysis",
                    f"Coded case {cid} needs at least one fully coded Commission-analysis passage.",
                    "Cases",
                    self.cases[cid][0],
                    "case_id",
                )
        if not self.counts["cases"]:
            self.missing(
                "empty_workbook",
                ["populated Cases records; examples and placeholders do not count"],
                "Cases",
            )
        if self.scope != "pilot":
            return
        states = Counter()
        for cid, (_, values) in self.cases.items():
            written = _date(values["order_date"])
            if (
                PILOT_ID.fullmatch(cid)
                and written
                and 2019 <= written.year <= 2026
                and values["state"] in SCHEMA["lists"]["State"]
                and values["sample_status"] in SCHEMA["lists"]["SampleStatus"]
            ):
                states[values["state"]] += 1
                self.sample[
                    "pre"
                    if written.year <= 2021
                    else "transition"
                    if written.year == 2022
                    else "post"
                ] += 1
        self.sample["states"] = dict(sorted(states.items()))
        for state, count in sorted(states.items()):
            if count > 4:
                self.issue(
                    "state_cap",
                    f"{state} has {count} eligible entries; maximum is four.",
                    "Cases",
                    severity="warning",
                )
        if self.counts["cases"] >= 15:
            self.sample["quotas_applied"] = True
            for period in ("pre", "post"):
                if self.sample[period] < 5:
                    self.issue(
                        "quota_" + period,
                        f"Need at least five {period} orders; counted {self.sample[period]}. 2022 counts toward neither quota.",
                        "Cases",
                        severity="warning",
                    )
            if states["VT"] < 2:
                self.issue(
                    "quota_vermont",
                    "Need at least two eligible Vermont orders.",
                    "Cases",
                    severity="warning",
                )
            if set(states) != set(SCHEMA["lists"]["State"]):
                self.issue(
                    "quota_states",
                    "Pilot does not cover all six New England states.",
                    "Cases",
                    severity="warning",
                )
            if self.counts["cases"] != 15:
                self.issue(
                    "pilot_size",
                    "Pilot requires exactly fifteen distinct populated orders.",
                    "Cases",
                    severity="warning",
                )
            required_nh = any(
                v["state"] == "NH"
                and re.fullmatch(r"(?:DE[ -]*)?24-070", str(v["docket"]).strip(), re.I)
                and v["sample_status"] == "Included by design"
                for _, v in self.cases.values()
            )
            if not required_nh:
                self.issue(
                    "required_nh_case",
                    "NH DE 24-070 must be included and flagged Included by design.",
                    "Cases",
                    severity="warning",
                )
        else:
            self.issue(
                "sample_in_progress",
                f"{self.counts['cases']}/15 populated case rows; full-sample quota checks deferred.",
                "Cases",
                severity="incomplete",
            )


def validate_sheet(path, *, mode="draft", scope="pilot"):
    """Return deterministic issues with cell addresses and counts without saving.

    Formula text is inspected, not executed; dates/adjusted ROE are checked in Python.
    """
    if mode not in ("draft", "complete") or scope not in ("pilot", "practice"):
        raise ValidationInputError("mode must be draft/complete and scope must be pilot/practice.")
    path = Path(path).expanduser().resolve()
    try:
        content = path.read_bytes()
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValidationInputError("Workbook exceeds the 100 MiB uncompressed input limit.")
        workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    except (
        OSError,
        BadZipFile,
        InvalidFileException,
        ValueError,
        KeyError,
        TypeError,
        SyntaxError,
    ) as error:
        raise ValidationInputError(f"Cannot read workbook: {error}") from error
    try:
        check = _Checks(workbook, mode, scope)
        if check.schema():
            check.case_rows()
            check.passage_rows()
            check.reliability_rows()
            check.cross_checks()
        counts = Counter(i["severity"] for i in check.issues)
        status = (
            "failed"
            if counts["error"]
            else "incomplete"
            if counts["incomplete"]
            else "passed_with_warnings"
            if counts["warning"]
            else "passed"
        )
        return dict(
            validator_version=VERSION,
            workbook_schema=SCHEMA["version"],
            schema_sha256=hashlib.sha256(SCHEMA_PATH.read_bytes()).hexdigest(),
            source_path=str(path),
            source_sha256=hashlib.sha256(content).hexdigest(),
            config=dict(mode=mode, scope=scope),
            status=status,
            errors=counts["error"],
            warnings=counts["warning"],
            incomplete=counts["incomplete"],
            counts=check.counts,
            sample=check.sample,
            derived_checks=check.derived,
            issues=check.issues,
            limits=[
                "Read-only validation; no values, categories, or missing facts repaired.",
                "Source accuracy, eligibility, note adequacy and reasoning labels are not certified.",
                "Excel formulas are not executed; v0.1.2 Cases formula text is checked on populated non-example rows.",
                "Period counts use written dates, not cached formula results; example rows are excluded only if their inputs match the frozen example.",
                "Kappa and aggregate sheet calculations remain M7/M8 work.",
            ],
        )
    finally:
        workbook.close()


def format_report(result):
    lines = [
        f"M6 {result['status']}: {result['errors']} errors, {result['warnings']} warnings, {result['incomplete']} unfinished items.",
        f"Workbook: {result['source_path']}",
        f"Scope: {result['config']['scope']}; mode: {result['config']['mode']}",
        f"Populated cases: {result['counts']['cases']}; passages: {result['counts']['passages']}; coded cases: {result['counts']['coded_cases']}.",
    ]
    for issue in result["issues"]:
        location = issue["sheet"] or "Workbook"
        if issue["cell"]:
            location += "!" + issue["cell"]
        elif issue["row"]:
            location += f" row {issue['row']}"
        lines.append(
            f"{issue['severity'].upper()} {location} [{issue['code']}]: {issue['message']}"
        )
    lines.append(
        "Validation does not establish source accuracy, research readiness, or independent coding reliability."
    )
    return "\n".join(lines) + "\n"


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--mode", choices=("draft", "complete"), default="draft")
    parser.add_argument("--scope", choices=("pilot", "practice"), default="pilot")
    parser.add_argument("--format", choices=("text", "json"), default="text")
    parser.add_argument(
        "--fail-on-warnings", action="store_true", help="Also fail on warnings or unfinished items."
    )
    args = parser.parse_args(argv)
    try:
        result = validate_sheet(args.workbook, mode=args.mode, scope=args.scope)
    except ValidationInputError as error:
        print(f"error: {error}", file=sys.stderr)
        return 2
    print(
        json.dumps(result, indent=2, ensure_ascii=False)
        if args.format == "json"
        else format_report(result),
        end="\n" if args.format == "json" else "",
    )
    return int(
        bool(
            result["errors"]
            or (args.fail_on_warnings and (result["warnings"] or result["incomplete"]))
        )
    )


if __name__ == "__main__":
    raise SystemExit(main())
