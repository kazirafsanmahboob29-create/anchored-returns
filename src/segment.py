"""M4: draft passages from an explicitly reviewed range, into a workbook copy.

Native, single-column, indented paragraphs are supported. Geometry is only a
boundary heuristic; the PDF's text must exactly match M2. No research labels.
"""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import sys
import tempfile

import pymupdf
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from src.locate_roe import _read_pages, _unique_object

HEADERS = [
    "passage_id",
    "case_id",
    "page_para",
    "section",
    "excerpt",
    "MKT",
    "method",
    "PEER",
    "OWN",
    "GRAD",
    "LAW",
    "CAP",
    "RISK",
    "AFF",
    "PERF",
    "DEF",
    "CONC",
    "primary_code",
    "stance",
    "coder",
    "notes",
]
PROTECTED = {"raw_orders", "practice_orders", "coding", "project_docs", "codebook", "literature"}


class SegmentationError(ValueError):
    """Invalid or unsupported input; never invent a plausible substitute."""


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _plan(path, pages):
    plan = json.loads(Path(path).read_text(), object_pairs_hook=_unique_object)
    required = {
        "case_id",
        "start_page",
        "end_page",
        "reviewer",
        "reviewer_type",
        "text_sha256",
        "pdf_sha256",
        "layout",
    }
    optional = {
        "start_at",
        "end_before",
        "dehyphenate_words",
        "paragraph_breaks",
        "footnote_continuations",
        "review_notes",
    }
    if (
        not isinstance(plan, dict)
        or not required <= plan.keys()
        or plan.keys() - required - optional
    ):
        raise SegmentationError("Unexpected or missing review-plan fields.")
    if not isinstance(plan["case_id"], str) or not re.fullmatch(
        r"[A-Za-z0-9][A-Za-z0-9_.-]*", plan["case_id"]
    ):
        raise SegmentationError("Invalid case_id.")
    if (
        plan["reviewer_type"] not in ("assistant", "human")
        or not isinstance(plan["reviewer"], str)
        or not plan["reviewer"].strip()
    ):
        raise SegmentationError("Identify the range reviewer and reviewer_type.")
    start, end = plan["start_page"], plan["end_page"]
    if type(start) is not int or type(end) is not int or not 1 <= start <= end <= len(pages):
        raise SegmentationError("Invalid inclusive viewer-page range.")
    layout = plan["layout"]
    if not isinstance(layout, dict) or set(layout) != {
        "body_left",
        "paragraph_indent",
        "header_bottom",
        "footer_top",
        "header_patterns",
    }:
        raise SegmentationError("Supply all layout fields explicitly.")
    for key in ("body_left", "paragraph_indent", "header_bottom", "footer_top"):
        if (
            type(layout[key]) not in (int, float)
            or not math.isfinite(layout[key])
            or layout[key] < 0
        ):
            raise SegmentationError("Layout coordinates must be finite nonnegative numbers.")
    if layout["paragraph_indent"] <= 0 or layout["header_bottom"] >= layout["footer_top"]:
        raise SegmentationError("Invalid layout geometry.")
    if not isinstance(layout["header_patterns"], list) or not all(
        isinstance(p, str) for p in layout["header_patterns"]
    ):
        raise SegmentationError("header_patterns must be a list of full-line regular expressions.")
    for key in ("start_at", "end_before"):
        if key in plan and (not isinstance(plan[key], str) or not plan[key].strip()):
            raise SegmentationError("Boundary anchors must be nonempty exact trimmed lines.")
    for key in ("paragraph_breaks", "footnote_continuations", "dehyphenate_words"):
        if not isinstance(plan.get(key, []), list) or not all(
            isinstance(v, str) and v for v in plan.get(key, [])
        ):
            raise SegmentationError(f"{key} must be a list of nonempty strings.")
    return plan


def _join(lines, allowed):
    result, changes = "", []
    for line in lines:
        value = line["text"]
        if not result:
            result = value
            continue
        if result.endswith(("-", "\u00ad")):
            before = re.search(r"([\w]+)[-\u00ad]$", result)
            after = re.match(r"([\w]+)", value)
            word = before[1] + after[1] if before and after else None
            remove = result.endswith("\u00ad") or word in allowed
            changes.append(
                dict(
                    before_line=line["id"],
                    operation="remove_line_end_hyphen" if remove else "keep_hyphen_join",
                    joined_word=word,
                )
            )
            result = (result[:-1] if remove else result) + value
        else:
            result += " " + value
    return result, changes


def segment(text_json, pdf, review_plan):
    """Return draft rows and an exhaustive nonblank source-line ledger; no writes."""
    text_json, pdf, review_plan = map(Path, (text_json, pdf, review_plan))
    pages = _read_pages(text_json.read_bytes())
    plan = _plan(review_plan, pages)
    if digest(text_json) != plan["text_sha256"] or digest(pdf) != plan["pdf_sha256"]:
        raise SegmentationError("Source fingerprint mismatch; review the changed inputs.")
    cfg = plan["layout"]
    try:
        patterns = [re.compile(p) for p in cfg["header_patterns"]]
    except re.error as error:
        raise SegmentationError(f"Invalid header pattern: {error}") from error
    ledger, bodies, notes = [], [], []
    used_overrides = set()
    with pymupdf.open(pdf) as doc:
        if len(doc) != len(pages) or doc.is_repaired or doc.is_encrypted:
            raise SegmentationError("PDF structure does not match the native extraction.")
        for number in range(plan["start_page"], plan["end_page"] + 1):
            page, saved = doc[number - 1], pages[number - 1]
            if saved["likely_scanned"] or page.get_text("text", sort=False) != saved["text"]:
                raise SegmentationError(
                    f"Page {number}: requires an exact, non-scanned native M2 extraction."
                )
            raw = []
            for block in page.get_text("dict", sort=False)["blocks"]:
                for line in block.get("lines", []):
                    value = "".join(span["text"] for span in line["spans"]).strip()
                    if value:
                        raw.append(
                            dict(
                                text=value,
                                bbox=list(line["bbox"]),
                                bold=any("bold" in span["font"].lower() for span in line["spans"]),
                            )
                        )
            source_lines = [
                (i, s.strip()) for i, s in enumerate(saved["text"].splitlines(), 1) if s.strip()
            ]
            if [v for _, v in source_lines] != [r["text"] for r in raw]:
                raise SegmentationError(
                    f"Page {number}: PDF layout/text line mapping is ambiguous."
                )
            for row, (index, _) in zip(raw, source_lines):
                row.update(id=f"{number}:{index}", page=number, line=index)
                row["disposition"] = "pending"
                ledger.append(row)
            kept = []
            for row in raw:
                y = row["bbox"][1]
                if (y < cfg["header_bottom"] or y > cfg["footer_top"]) and any(
                    p.fullmatch(row["text"]) for p in patterns
                ):
                    row["disposition"] = "running_header_or_footer"
                else:
                    kept.append(row)
            # Numeric markers must align with a separate indented text line.
            fn, body = [], []
            for i, row in enumerate(kept):
                nxt = kept[i + 1] if i + 1 < len(kept) else None
                marker = (
                    re.fullmatch(r"\d+", row["text"])
                    and nxt
                    and abs(row["bbox"][1] - nxt["bbox"][1]) <= 3
                    and abs(row["bbox"][0] - cfg["body_left"]) <= 4
                    and nxt["bbox"][0] >= cfg["body_left"] + cfg["paragraph_indent"] - 4
                )
                continuation = row["id"] in plan.get("footnote_continuations", [])
                if continuation:
                    used_overrides.add(row["id"])
                if continuation:
                    if not notes or notes[-1]["lines"][-1]["page"] != number - 1:
                        raise SegmentationError(
                            "Footnote continuation needs a note on the preceding page."
                        )
                    fn = notes[-1]["lines"]
                elif marker:
                    fn = []
                    notes.append(dict(kind="footnote", marker=row["text"], lines=fn))
                if fn or marker or continuation:
                    fn.append(row)
                else:
                    body.append(row)
            # Anchors affect body only: footnotes on the last page survive clipping.
            for key, active, before in [
                ("start_at", number == plan["start_page"], True),
                ("end_before", number == plan["end_page"], False),
            ]:
                if key in plan and active:
                    matches = [i for i, row in enumerate(body) if row["text"] == plan[key]]
                    if len(matches) != 1:
                        raise SegmentationError(
                            f"{key} must match exactly one body line on page {number}."
                        )
                    cut = matches[0]
                    removed, body = (body[:cut], body[cut:]) if before else (body[cut:], body[:cut])
                    for row in removed:
                        row["disposition"] = "outside_reviewed_body_range"
            bodies.extend(body)
    groups, current = [], None
    for i, row in enumerate(bodies):
        prev = bodies[i - 1] if i else None
        nxt = bodies[i + 1] if i + 1 < len(bodies) else None
        heading = row["bold"] or bool(
            re.fullmatch(r"(?:[A-Za-z]|[ivxlcdmIVXLCDM]+|\d+)\.", row["text"])
        )
        same_line = (
            prev is not None
            and prev["page"] == row["page"]
            and abs(prev["bbox"][1] - row["bbox"][1]) <= 3
        )
        indent = abs(row["bbox"][0] - (cfg["body_left"] + cfg["paragraph_indent"])) <= 4
        # A first-line indent followed by a left-margin line distinguishes a new
        # paragraph from an indented quotation. Uncertain layout remains a draft.
        first_indent = indent and (nxt is None or abs(nxt["bbox"][0] - cfg["body_left"]) <= 4)
        manual = row["id"] in plan.get("paragraph_breaks", [])
        if manual:
            used_overrides.add(row["id"])
        if (
            current is None
            or manual
            or (not same_line and (heading or current["kind"] == "heading" or first_indent))
        ):
            current = dict(kind="heading" if heading else "body", lines=[])
            groups.append(current)
        current["lines"].append(row)
    overrides = set(plan.get("paragraph_breaks", []) + plan.get("footnote_continuations", []))
    if overrides != used_overrides:
        raise SegmentationError(
            f"Unused or invalid line overrides: {sorted(overrides - used_overrides)}"
        )
    # Separate footnote passages; stable source order places a note after the
    # body passage that starts before it, even when that body spans a page break.
    groups.extend(notes)
    groups.sort(key=lambda group: (group["lines"][0]["page"], group["lines"][0]["line"]))
    rows, counters = [], Counter()
    for index, group in enumerate(groups, 1):
        lines = group["lines"]
        passage_id = f"{plan['case_id']}-{index:04d}"
        text, changes = _join(lines, set(plan.get("dehyphenate_words", [])))
        refs = []
        for number in dict.fromkeys(row["page"] for row in lines):
            counters[number] += 1
            refs.append(f"PDF p.{number}/part {counters[number]}")
        message = (
            f"DRAFT; {group['kind']}; {plan['reviewer_type']}-reviewed range; check PDF boundaries"
        )
        if group["kind"] == "footnote":
            message += f"; footnote {group['marker']}"
        if any(c["operation"] == "keep_hyphen_join" for c in changes):
            message += "; check retained line-end hyphen(s)"
        if group["kind"] == "footnote" and not re.search(r"[.!?][”\")]*$", text):
            message += "; possible continuation: inspect next page"
        row = dict(
            passage_id=passage_id,
            case_id=plan["case_id"],
            page_para="; ".join(refs),
            excerpt=text,
            notes=message,
            kind=group["kind"],
            source_line_ids=[r["id"] for r in lines],
            normalization=changes,
        )
        rows.append(row)
        for line in lines:
            if line["disposition"] != "pending":
                raise SegmentationError("Source line assigned more than once.")
            line["disposition"] = passage_id
    if not rows or any(row["disposition"] == "pending" for row in ledger):
        raise SegmentationError("Empty output or unaccounted source lines.")
    return dict(
        algorithm_version="1.0",
        status="draft_requires_pdf_review",
        review_plan=plan,
        text_sha256=digest(text_json),
        pdf_sha256=digest(pdf),
        review_plan_sha256=digest(review_plan),
        page_numbering="one-based PDF viewer positions; part numbers are draft units, not printed paragraph numbers",
        passages=rows,
        source_line_ledger=ledger,
        limits=[
            "Native, single-column indented prose only; geometry is a heuristic.",
            "Review quotations, tables, footnote continuations and all passage boundaries.",
            "No section, code, stance, coder, or case-level numeric fields filled.",
        ],
    )


def write_draft(text_json, pdf, review_plan, template, output_dir):
    """Publish a fresh bundle. Refuse existing outputs and nonempty coding sheets."""
    output = Path(output_dir).expanduser()
    resolved = output.resolve()
    if output.exists() or output.is_symlink() or any(p in PROTECTED for p in resolved.parts):
        raise SegmentationError(
            "Use a new output directory outside protected source/coding folders."
        )
    for parent in [output, *output.parents]:
        if parent.is_symlink():
            raise SegmentationError("Output path must not contain symlinks.")
    result = segment(text_json, pdf, review_plan)
    template = Path(template)
    template_hash = digest(template)
    wb = load_workbook(template)
    if "Passages" not in wb.sheetnames or [c.value for c in wb["Passages"][1]] != HEADERS:
        raise SegmentationError("Workbook Passages schema differs from v0.1.2.")
    sheet = wb["Passages"]
    if any(cell.value is not None for row in sheet.iter_rows(min_row=3) for cell in row):
        raise SegmentationError("Use an empty template copy; Passages already contains data.")
    if len(result["passages"]) + 2 > sheet.max_row:
        raise SegmentationError(
            "Draft exceeds template capacity; extend reviewed validation ranges first."
        )
    for number, row in enumerate(result["passages"], 3):
        for field in ("passage_id", "case_id", "page_para", "excerpt", "notes"):
            value = row[field]
            if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
                raise SegmentationError(
                    "Excerpt cannot be represented losslessly in an Excel cell."
                )
            cell = sheet.cell(number, HEADERS.index(field) + 1)
            cell.value = value
            cell.data_type = "s"  # Source text beginning with '=' is not executable.
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[number].height = min(
            409, max(70, math.ceil(len(row["excerpt"]) / 80) * 15)
        )
    sheet.freeze_panes = "F3"
    sheet.sheet_view.zoomScale = 85
    wb.active = wb.sheetnames.index("Passages")
    # Blinding requested by the protocol. Formulas and input values are preserved.
    for col in range(42, 58):
        from openpyxl.utils import get_column_letter

        wb["Cases"].column_dimensions[get_column_letter(col)].hidden = True
    for name in ("Case_Codes", "Summary"):
        wb[name].sheet_state = "hidden"
    result.update(
        template_sha256=template_hash,
        workbook="passages_draft.xlsx",
        workbook_note="Practice/review copy only; original example row and Cases placeholders retained. Outcome summaries hidden.",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    # The bundle is built in a unique sibling; reserve destination before publishing
    # so a concurrent run cannot replace an existing bundle, even an empty folder.
    temporary = Path(tempfile.mkdtemp(prefix=".m4-", dir=output.parent))
    reserved = False
    try:
        wb.save(temporary / result["workbook"])
        result["workbook_sha256"] = digest(temporary / result["workbook"])
        (temporary / "passages.json").write_text(
            json.dumps(result["passages"], indent=2, ensure_ascii=False) + "\n"
        )
        result["passages_sha256"] = digest(temporary / "passages.json")
        (temporary / "audit.json").write_text(
            json.dumps(result, indent=2, ensure_ascii=False) + "\n"
        )
        if digest(template) != template_hash:
            raise SegmentationError("Template changed during export.")
        output.mkdir()
        reserved = True
        for path in temporary.iterdir():
            path.rename(output / path.name)
    except Exception:
        if reserved:
            shutil.rmtree(output)
        raise
    finally:
        shutil.rmtree(temporary)
    return dict(
        output_dir=str(resolved),
        n_passages=len(result["passages"]),
        kinds=dict(Counter(row["kind"] for row in result["passages"])),
        status=result["status"],
    )


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    for argument in ("text_json", "pdf", "review_plan", "template", "output_dir"):
        parser.add_argument(argument, type=Path)
    try:
        result = write_draft(**vars(parser.parse_args(argv)))
    except (ValueError, OSError, KeyError, TypeError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 2
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
