"""M8: regenerate descriptive tables, figures and a memo from one coding workbook.

Inputs remain untouched. Missing observations are explicit; no classification,
source acquisition, imputation, or causal inference is performed by this module.
"""

from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
from decimal import Decimal
import hashlib
import html
from io import BytesIO
import json
from pathlib import Path
import shutil
import statistics
import sys
import textwrap

import matplotlib

matplotlib.use("Agg")
from matplotlib import pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook

from src.validate_sheet import (
    validate_sheet,
    ValidationInputError,
    SCHEMA,
    CODES,
    ROE_FIELDS,
    _blank,
    _date,
)
from src.reliability import reliability, ReliabilityInputError, ReliabilityMismatchError

VERSION = "1.0"
REGIMES = (
    "Pre (2019-21)",
    "Transition (2022)",
    "Post (2023-26)",
    "Practice / out of frame",
    "Date missing",
)
NUMBER_FIELDS = (
    "advocate_rec",
    "staff_rec",
    "alj_rec",
    "requested_roe_initial",
    "requested_roe",
    "awarded_roe",
    "awarded_roe_final",
    "prior_authorized_roe",
    "nat_avg_benchmark",
)
NUMBER_LABELS = (
    "Advocate",
    "Staff",
    "ALJ",
    "Initial request",
    "Final request",
    "Base award",
    "Final award",
    "Prior award",
    "National benchmark",
)
RECOVERY_FIELDS = (
    *ROE_FIELDS,
    "perf_adjustment_bp",
    "ust10_at_order",
    "ust10_at_prior_order",
    "equity_ratio_awarded",
)
REVIEW_HEADERS = ["case_id", "top_hit", "reviewer", "reviewer_type", "reviewed_on", "evidence"]
PROTECTED = {
    "raw_orders",
    "practice_orders",
    "coding",
    "project_docs",
    "codebook",
    "literature",
    ".git",
    ".conda",
    ".venv",
}
COLORS = (
    "#76528b",
    "#238488",
    "#627283",
    "#af7628",
    "#d15a36",
    "#194f7b",
    "#161e2c",
    "#6f8547",
    "#b65073",
)
MARKERS = ("o", "s", "^", "v", "D", "P", "X", "<", ">")
METHODS = [
    "All figures are descriptive. Similarity to a reference does not establish anchoring or causation.",
    "Number-line inputs use ROE fractions; display values are percentages. Final award requires an explicit signed basis-point adjustment, including zero.",
    "Scatter uses current BASE ROE minus prior authorized ROE and Treasury change over the same recorded order dates, both in percentage points. No regression is fitted.",
    "Shares use only fully coded Commission-analysis passages. They are pooled passage shares, not equal-weight order means; co-occurring codes can sum above 100%.",
    "Code presence includes Relied on, Considered and Rejected passages. It does not mean endorsement, and primary-code stance is not assigned to secondary codes.",
    "Expected earnings moves between market and benchmark groups only in the sensitivity table. PEER/OWN/GRAD inputs are not rewritten. Mixed or unspecified market methods produce bounds.",
    "Recovery means a numeric value was entered with case source notes, not that every absent field should exist. Missing, not applicable and not yet collected are not distinguishable in this schema.",
    "M3 hit rates use only explicit M3_Review records, separately for assistant and human reviewers. A missing record is not a failed hit.",
    "Reliability compares saved values through M7. Rater identities, blinding and independence cannot be established from the four rating columns.",
]


class ReportInputError(ValueError):
    """Unsafe destination, invalid workbook, or inconsistent upstream results."""


def _require(condition, message):
    if not condition:
        raise ReportInputError(message)


def _hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def regime(value):
    d = _date(value)
    if d is None:
        return REGIMES[4]
    if not 2019 <= d.year <= 2026:
        return REGIMES[3]
    return REGIMES[0] if d.year <= 2021 else REGIMES[1] if d.year == 2022 else REGIMES[2]


def _rows(book, sheet, headers):
    for row in range(2, book[sheet].max_row + 1):
        values = {key: book[sheet].cell(row, col).value for col, key in enumerate(headers, 1)}
        if not all(_blank(v) for v in values.values()):
            yield row, values


def _reviews(book, case_ids):
    """Read an optional, explicit review log without inferring hits from keywords."""
    if "M3_Review" not in book:
        return []
    ws = book["M3_Review"]
    _require(
        ws.max_row <= 10000 and ws.max_column <= len(REVIEW_HEADERS),
        "M3_Review: unexpected dimensions.",
    )
    _require(
        [c.value for c in ws[1]] == REVIEW_HEADERS,
        "M3_Review: expected headers " + ", ".join(REVIEW_HEADERS),
    )
    result = []
    seen = set()
    for row, record in _rows(book, "M3_Review", REVIEW_HEADERS):
        _require(
            not any(ws.cell(row, c).data_type in ("f", "e") for c in range(1, 7)),
            f"M3_Review row {row}: input formula/error.",
        )
        _require(record["case_id"] in case_ids, f"M3_Review row {row}: unknown case.")
        _require(
            type(record["top_hit"]) in (int, float) and record["top_hit"] in (0, 1),
            f"M3_Review row {row}: top_hit must be numeric 0/1.",
        )
        _require(
            record["reviewer_type"] in ("human", "assistant"),
            f"M3_Review row {row}: identify reviewer type.",
        )
        _require(
            all(isinstance(record[k], str) and record[k].strip() for k in ("reviewer", "evidence")),
            f"M3_Review row {row}: reviewer/evidence required.",
        )
        d = _date(record["reviewed_on"])
        _require(d is not None, f"M3_Review row {row}: invalid date.")
        key = (record["case_id"], record["reviewer_type"])
        _require(key not in seen, f"M3_Review row {row}: duplicate case/reviewer-type review.")
        seen.add(key)
        result.append({**record, "reviewed_on": d.isoformat(), "source_row": row})
    return result


def _complete_passage(p):
    fields = (*CODES, "section", "primary_code", "stance", "coder", "excerpt", "page_para")
    return all(not _blank(p[f]) for f in fields) and (p["MKT"] != 1 or not _blank(p["method"]))


def _shares(passages, groups):
    shares = []
    sensitivity = []
    for group, ids in groups:
        rows = [p for p in passages if p["case_id"] in ids]
        n = len(rows)
        for code in CODES:
            count = sum(p[code] == 1 for p in rows)
            shares.append(
                dict(
                    group=group,
                    code=code,
                    count=count,
                    denominator=n,
                    share=count / n if n else None,
                )
            )
        market = [p for p in rows if p["MKT"] == 1]
        anchors = {p["passage_id"] for p in rows if any(p[k] == 1 for k in ("PEER", "OWN", "GRAD"))}
        ee = {p["passage_id"] for p in market if p["method"] == "Expected earnings"}
        ambiguous = {
            p["passage_id"] for p in market if p["method"] in ("Multiple", "Unspecified", "n.a.")
        }
        for variant in ("EE as market", "EE as benchmark"):
            low_m = high_m = len(market)
            low_b = high_b = len(anchors)
            if variant == "EE as benchmark":
                high_m -= len(ee)
                low_m = high_m - len(ambiguous)
                low_b = len(anchors | ee)
                high_b = len(anchors | ee | ambiguous)
            sensitivity.append(
                dict(
                    group=group,
                    variant=variant,
                    denominator=n,
                    expected_earnings=len(ee),
                    ambiguous_method=len(ambiguous),
                    market_min=low_m,
                    market_max=high_m,
                    benchmark_min=low_b,
                    benchmark_max=high_b,
                    market_share_min=low_m / n if n else None,
                    market_share_max=high_m / n if n else None,
                    benchmark_share_min=low_b / n if n else None,
                    benchmark_share_max=high_b / n if n else None,
                )
            )
    return shares, sensitivity


def analyse_workbook(path, *, scope="pilot", synthetic=False):
    """Validate and derive all report data from one unchanged workbook.

    Recompute arithmetic from typed inputs, retaining missingness and provenance.
    """
    path = Path(path).expanduser().resolve()
    try:
        validation = validate_sheet(path, scope=scope)
        bad = [i for i in validation["issues"] if i["severity"] == "error"]
        _require(not bad, "M6 validation errors: " + json.dumps(bad[:5]))
        _require(
            not any(i["code"] == "missing_numeric_sources" for i in validation["issues"]),
            "Numeric sources_notes are required before reporting entered numbers.",
        )
        rel = reliability(path)
        _require(
            rel["source_sha256"] == validation["source_sha256"] == _hash(path),
            "Input changed during validation.",
        )
        book = load_workbook(BytesIO(path.read_bytes()), data_only=False, keep_links=False)
    except (ValidationInputError, ReliabilityInputError, ReliabilityMismatchError, OSError) as e:
        raise ReportInputError(str(e)) from e
    try:
        cases = []
        for row, c in _rows(book, "Cases", SCHEMA["case_headers"][:41]):
            if c["case_id"] == "EX-01" or all(_blank(v) for k, v in c.items() if k != "case_id"):
                continue
            _require(
                synthetic
                or not str(c["sources_notes"] or "").upper().startswith("SYNTHETIC TEST ONLY"),
                "Synthetic fixture requires --synthetic.",
            )
            typed_fields = set(RECOVERY_FIELDS) | {
                "minutes_spent",
                "order_date",
                "prior_order_date",
                "open_meeting_date",
                "date_coded",
            }
            _require(
                all(_blank(v) or isinstance(v, str) for k, v in c.items() if k not in typed_fields),
                f"Cases row {row}: metadata and source notes must be text.",
            )
            c = {k: (None if _blank(v) else v) for k, v in c.items()}
            for key in ("order_date", "prior_order_date", "open_meeting_date", "date_coded"):
                d = _date(c[key])
                c[key] = d.isoformat() if d else None
            c["source_row"] = row
            c["regime"] = regime(c["order_date"])
            base, bp = c["awarded_roe"], c["perf_adjustment_bp"]
            c["awarded_roe_final"] = (
                float(Decimal(str(base)) + Decimal(str(bp)) / 10000)
                if base is not None and bp is not None
                else None
            )
            cases.append(c)
        cases.sort(key=lambda c: (REGIMES.index(c["regime"]), c["order_date"] or "", c["case_id"]))
        ids = {c["case_id"] for c in cases}
        passages = [
            {**p, "source_row": row}
            for row, p in _rows(book, "Passages", SCHEMA["passage_headers"])
            if p["case_id"] != "EX-01"
        ]
        commissioned = [p for p in passages if p["section"] == "Commission analysis"]
        complete = [p for p in commissioned if _complete_passage(p)]
        points = []
        excluded = []
        required = (
            "awarded_roe",
            "prior_authorized_roe",
            "ust10_at_order",
            "ust10_at_prior_order",
            "order_date",
            "prior_order_date",
        )
        for c in cases:
            missing = [f for f in required if c[f] is None]
            if missing:
                excluded.append(dict(case_id=c["case_id"], missing_fields=missing))
                continue
            diff = lambda a, b: float((Decimal(str(c[a])) - Decimal(str(c[b]))) * 100)
            points.append(
                dict(
                    case_id=c["case_id"],
                    state=c["state"],
                    regime=c["regime"],
                    delta_roe_pp=diff("awarded_roe", "prior_authorized_roe"),
                    delta_ust10_pp=diff("ust10_at_order", "ust10_at_prior_order"),
                    order_date=c["order_date"],
                    prior_order_date=c["prior_order_date"],
                )
            )
        groups = (
            [("All", ids)]
            + [
                (r, {c["case_id"] for c in cases if c["regime"] == r})
                for r in REGIMES
                if any(c["regime"] == r for c in cases)
            ]
            + [(c["case_id"], {c["case_id"]}) for c in cases]
        )
        shares, sensitivity = _shares(complete, groups)
        reviews = _reviews(book, ids)
        hits = []
        for kind in ("assistant", "human"):
            records = [r for r in reviews if r["reviewer_type"] == kind]
            n = len(records)
            count = sum(r["top_hit"] for r in records)
            hits.append(
                dict(
                    reviewer_type=kind,
                    hits=count,
                    reviewed=n,
                    rate=count / n if n else None,
                    unreviewed=len(cases) - n,
                )
            )
        coverage = [
            dict(
                field=f,
                entered=sum(c[f] is not None for c in cases),
                orders=len(cases),
                share=sum(c[f] is not None for c in cases) / len(cases) if cases else None,
            )
            for f in RECOVERY_FIELDS
        ]
        feasibility = []
        for c in cases:
            ps = [p for p in passages if p["case_id"] == c["case_id"]]
            feasibility.append(
                dict(
                    case_id=c["case_id"],
                    state=c["state"],
                    regime=c["regime"],
                    minutes_spent=c["minutes_spent"],
                    numeric_fields_entered=sum(c[f] is not None for f in RECOVERY_FIELDS),
                    numeric_fields_tracked=len(RECOVERY_FIELDS),
                    roe_section_recorded=c["roe_section_pages"] is not None,
                    n_passages=len(ps),
                    n_commission=sum(p["section"] == "Commission analysis" for p in ps),
                    n_complete_commission=sum(p["case_id"] == c["case_id"] for p in complete),
                )
            )
        minutes = [c["minutes_spent"] for c in cases if c["minutes_spent"] is not None]
        result = dict(
            module="M8",
            version=VERSION,
            scope=scope,
            dataset_kind="synthetic" if synthetic else scope,
            source_path=str(path),
            source_sha256=validation["source_sha256"],
            n_cases=len(cases),
            cases=cases,
            validation=validation,
            reliability=rel,
            coding=dict(
                n_passages=len(passages),
                n_commission=len(commissioned),
                n_complete_commission=len(complete),
                n_incomplete_commission=len(commissioned) - len(complete),
                n_unassigned_section=sum(_blank(p["section"]) for p in passages),
            ),
            scatter=dict(
                n_plotted=len(points), n_excluded=len(excluded), points=points, excluded=excluded
            ),
            code_shares=shares,
            sensitivity=sensitivity,
            feasibility=dict(
                orders=feasibility,
                minutes_recorded=len(minutes),
                median_minutes=statistics.median(minutes) if minutes else None,
                m3_reviews=reviews,
                m3_hit_rates=hits,
                field_coverage=coverage,
            ),
            methods=METHODS,
        )
        _require(_hash(path) == result["source_sha256"], "Input changed during analysis.")
        return result
    finally:
        book.close()


def _banner(data):
    if data["dataset_kind"] == "synthetic":
        return "SYNTHETIC DEMONSTRATION - not research findings"
    if data["scope"] == "practice":
        return "PRACTICE / INCOMPLETE RESEARCH DATA"
    if data["validation"]["status"] != "passed":
        return "INCOMPLETE PILOT DATA - descriptive export"
    return "DESCRIPTIVE PILOT EXPORT - verify coverage and provenance"


def _style(ax):
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#dce2e7", linewidth=0.6)
    ax.set_axisbelow(True)


def _empty(ax, message):
    ax.text(
        0.5,
        0.5,
        message,
        ha="center",
        va="center",
        transform=ax.transAxes,
        color="#586575",
        wrap=True,
    )
    ax.set_xticks([])
    ax.set_yticks([])


def plot_rate_shock(data):
    """Use identical percentage-point scales so the reference line is truly 45 degrees."""
    fig, ax = plt.subplots(figsize=(6.4, 6.4))
    ps = data["scatter"]["points"]
    _style(ax)
    values = [0] + [p[k] for p in ps for k in ("delta_roe_pp", "delta_ust10_pp")]
    lo, hi = min(values), max(values)
    pad = max((hi - lo) * 0.12, 0.25)
    lo -= pad
    hi += pad
    ax.plot([lo, hi], [lo, hi], "--", color="#838b93", lw=1, label="One-for-one change")
    seen_regimes = set()
    for i, p in enumerate(
        sorted(ps, key=lambda p: (p["delta_ust10_pp"], p["delta_roe_pp"], p["case_id"]))
    ):
        label = p["regime"] if p["regime"] not in seen_regimes else "_nolegend_"
        seen_regimes.add(p["regime"])
        ax.scatter(
            p["delta_ust10_pp"],
            p["delta_roe_pp"],
            s=42,
            color=COLORS[REGIMES.index(p["regime"])],
            label=label,
            zorder=3,
        )
        direction = -1 if p["delta_roe_pp"] > (lo + hi) / 2 else 1
        ax.annotate(
            p["case_id"],
            (p["delta_ust10_pp"], p["delta_roe_pp"]),
            xytext=(5, direction * (14 + (i % 3) * 14)),
            textcoords="offset points",
            fontsize=8,
            arrowprops=dict(arrowstyle="-", lw=0.45, color="#88949e"),
        )
    ax.set(
        xlim=(lo, hi),
        ylim=(lo, hi),
        xlabel="Change in 10-year Treasury (percentage points)",
        ylabel="Change in base ROE (percentage points)",
        title=f"Rate-shock comparison: descriptive, n = {len(ps)}",
    )
    ax.set_aspect("equal", adjustable="box")
    ax.legend(loc="lower right", frameon=False, fontsize=8)
    if not ps:
        _empty(ax, "No complete prior/current ROE and Treasury intervals")
    fig.text(0.5, 0.02, _banner(data), ha="center", fontsize=8, color="#8b4c24")
    fig.tight_layout(rect=(0, 0.05, 1, 1))
    return fig


def plot_number_lines(data, case=None):
    rows = [case] if case else data["cases"]
    if case:
        fig, ax = plt.subplots(figsize=(7, 4.4))
        ys = list(range(9))
        for i, key in enumerate(NUMBER_FIELDS):
            v = case[key]
            if v is not None:
                ax.scatter(v * 100, i, marker=MARKERS[i], color=COLORS[i], s=62)
            ax.text(
                1.01,
                i,
                "missing" if v is None else f"{v * 100:.2f}%",
                transform=ax.get_yaxis_transform(),
                va="center",
                fontsize=9,
            )
        ax.set_yticks(ys, NUMBER_LABELS)
        ax.invert_yaxis()
        ax.set_title(case["case_id"] + " | " + case["regime"], fontsize=11)
    else:
        fig, ax = plt.subplots(figsize=(8, max(4.2, len(rows) * 0.36 + 2.4)))
        for j, key in enumerate(NUMBER_FIELDS):
            pts = [(i, c[key] * 100) for i, c in enumerate(rows) if c[key] is not None]
            ax.scatter(
                [v for _, v in pts],
                [i + (j - 4) * 0.055 for i, _ in pts],
                label=NUMBER_LABELS[j],
                marker=MARKERS[j],
                color=COLORS[j],
                s=34,
                zorder=3,
            )
        ax.set_yticks(
            range(len(rows)),
            [f"{c['case_id']}  {c['state'] or '?'}  {(c['order_date'] or '?')[:4]}" for c in rows],
            fontsize=9,
        )
        for i, c in enumerate(rows):
            ax.axhspan(
                i - 0.45,
                i + 0.45,
                color="#f1f4f6" if REGIMES.index(c["regime"]) % 2 == 0 else "#ffffff",
                zorder=0,
            )
            if i and c["regime"] != rows[i - 1]["regime"]:
                ax.axhline(i - 0.5, color="#9aa8b4", lw=1)
        ax.set_ylim(len(rows) - 0.5, -0.6)
        ax.set_title(
            "ROE comparisons by written-order regime\nPre 2019-21 | Transition 2022 | Post 2023-26",
            fontsize=11,
        )
        ax.legend(
            loc="upper center", bbox_to_anchor=(0.5, -0.16), ncol=3, fontsize=8, frameon=False
        )
        if not rows or not any(c[k] is not None for c in rows for k in NUMBER_FIELDS):
            _empty(ax, "No ROE values entered for comparison")
    all_values = [c[k] * 100 for c in data["cases"] for k in NUMBER_FIELDS if c[k] is not None]
    lo, hi = (min(all_values), max(all_values)) if all_values else (3, 20)
    pad = max((hi - lo) * 0.12, 0.15)
    ax.set_xlim(lo - pad, hi + pad)
    ax.set_xlabel("Return on equity (%)")
    _style(ax)
    fig.text(0.5, 0.02, _banner(data), ha="center", fontsize=8, color="#8b4c24")
    fig.tight_layout(rect=(0, 0.05, 0.91 if case else 1, 1))
    return fig


def plot_code_shares(data):
    rows = [r for r in data["code_shares"] if r["group"] == "All"]
    fig, ax = plt.subplots(figsize=(6.5, 4.6))
    n = rows[0]["denominator"]
    if n:
        ax.barh(range(len(rows)), [r["share"] * 100 for r in rows], color="#245c78", height=0.65)
        for i, r in enumerate(rows):
            ax.text(min(r["share"] * 100 + 2, 102), i, f"{r['count']}/{n}", va="center", fontsize=9)
        ax.set_yticks(range(len(rows)), [r["code"] for r in rows])
        ax.invert_yaxis()
    else:
        _empty(ax, "No fully coded Commission-analysis passages")
    ax.set_xlim(0, 115)
    ax.set_xticks([0, 25, 50, 75, 100])
    ax.set_xlabel("Share of fully coded Commission-analysis passages (%)")
    ax.set_title(f"Code presence (multiple codes allowed; n = {n})", fontsize=11)
    _style(ax)
    fig.text(0.5, 0.02, _banner(data), ha="center", fontsize=8, color="#8b4c24")
    fig.tight_layout(rect=(0, 0.05, 1, 1))
    return fig


def _display(v):
    if v is None:
        return "not available"
    if type(v) is float:
        return f"{v:.4g}"
    return str(v)


def _html_table(rows, columns):
    if not rows:
        return '<p class="missing">No observations recorded.</p>'
    esc = lambda v: html.escape(_display(v))
    head = "".join("<th>" + esc(label) + "</th>" for _, label in columns)
    body = "".join(
        "<tr>" + "".join("<td>" + esc(r.get(key)) + "</td>" for key, _ in columns) + "</tr>"
        for r in rows
    )
    return (
        '<div class="scroll"><table><thead><tr>'
        + head
        + "</tr></thead><tbody>"
        + body
        + "</tbody></table></div>"
    )


def _sensitivity_view(data):
    rows = []
    for r in data["sensitivity"]:
        if r["group"] != "All":
            continue

        def share(a, b):
            if not r["denominator"]:
                return "not available"
            lo, hi = r[a] / r["denominator"] * 100, r[b] / r["denominator"] * 100
            return f"{lo:.1f}%" if lo == hi else f"{lo:.1f}-{hi:.1f}%"

        rows.append(
            dict(
                variant=r["variant"],
                market=share("market_min", "market_max"),
                benchmark=share("benchmark_min", "benchmark_max"),
                ambiguous=r["ambiguous_method"],
            )
        )
    return rows


def _write_html(data, destination):
    esc = html.escape
    parts = [
        '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">',
        "<title>Anchored Returns | Pilot report</title><style>body{font:16px/1.55 system-ui,sans-serif;background:#f4f5f6;color:#172b3c;margin:0}main{max-width:1100px;margin:auto;padding:44px 30px;background:white}h1{font-size:36px;line-height:1.2}h2{margin-top:40px;border-top:1px solid #ccd4da;padding-top:20px}h3{font-size:19px}p{max-width:90ch}.notice{background:#fff3df;padding:16px;border-left:4px solid #a77337}.metrics{display:flex;gap:18px;flex-wrap:wrap}.metrics div{padding:16px 22px;background:#f0f4f7}.metrics strong{display:block;font-size:28px}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(360px,1fr));gap:20px}img{width:100%;height:auto}.scroll{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:14px}th{text-align:left;background:#e8eef2}td,th{padding:9px;border-bottom:1px solid #dee3e7}td{vertical-align:top;overflow-wrap:anywhere}.missing{color:#7f5424}details{margin-top:18px}a{color:#165b89}code{overflow-wrap:anywhere}@media print{body{background:white}main{padding:0}.scroll{overflow:visible}details{display:block}}</style><main>",
        "<p>ANCHORED RETURNS / M8</p><h1>Pilot reporting</h1>",
        '<p class="notice">' + esc(_banner(data)) + "</p>",
        f"<p>Workbook validation: {esc(data['validation']['status'])}; {data['validation']['warnings']} warnings and {data['validation']['incomplete']} unfinished items. Output generation does not complete the research dataset.</p>",
        f'<div class="metrics"><div><strong>{data["n_cases"]}</strong>populated orders</div><div><strong>{data["scatter"]["n_plotted"]}</strong>complete comparison intervals</div><div><strong>{data["coding"]["n_complete_commission"]}</strong>fully coded analysis passages</div></div>',
        "<p>These outputs describe the workbook supplied. They do not establish source accuracy, sample eligibility, independent coding, or an anchoring effect. Missing fields remain unavailable.</p>",
        '<p><a href="memo.pdf">Two-page memo</a> · <a href="analysis.json">Analysis data</a> · <a href="audit.json">Input and output fingerprints</a></p>',
        "<h2>ROE comparisons</h2><p>Common percentage scale, ordered by Pre, Transition, Post, then practice/missing dates. Unentered values are omitted. Final award requires an explicit adjustment; blank does not imply zero. A cited national figure is not substituted for the defined benchmark.</p>",
        '<img src="figures/number_lines.png" alt="Order-level ROE comparisons grouped by regime">',
        '<div class="grid"><div><h2>Rate-shock comparison</h2><img src="figures/rate_shock.png" alt="Change in base ROE versus Treasury change"><p>Both axes use percentage points and equal scaling. Each point is one complete order interval.</p></div><div><h2>Code presence</h2><img src="figures/code_shares.png" alt="Code presence shares"><p>Only fully coded Commission-analysis passages enter the denominator. Co-occurring codes are counted independently.</p></div></div>',
        f"<p>Coverage: {data['coding']['n_complete_commission']}/{data['coding']['n_commission']} identified Commission-analysis passages are fully coded; {data['coding']['n_unassigned_section']} passages have no section assignment.</p>",
        "<h2>Expected earnings sensitivity</h2><p>The benchmark group is the union of PEER, OWN and GRAD. In the second view it also includes expected earnings. Ranges reflect Multiple, Unspecified or n.a. market methods whose expected-earnings content is unknown.</p>",
        _html_table(
            _sensitivity_view(data),
            [
                ("variant", "Treatment"),
                ("market", "Market share"),
                ("benchmark", "Benchmark share"),
                ("ambiguous", "Ambiguous methods"),
            ],
        ),
        "<h2>Feasibility</h2><p>Median recorded minutes per order: "
        + esc(_display(data["feasibility"]["median_minutes"]))
        + f"; timing coverage {data['feasibility']['minutes_recorded']}/{data['n_cases']}. Numeric recovery counts entered values, not source certification.</p>",
        _html_table(
            data["feasibility"]["orders"],
            [
                ("case_id", "Order"),
                ("regime", "Regime"),
                ("minutes_spent", "Minutes"),
                ("numeric_fields_entered", "Numeric fields entered"),
                ("numeric_fields_tracked", "Fields tracked"),
                ("n_complete_commission", "Coded analysis passages"),
            ],
        ),
        _html_table(
            data["feasibility"]["m3_hit_rates"],
            [
                ("reviewer_type", "M3 reviewer type"),
                ("hits", "Top hits"),
                ("reviewed", "Reviewed"),
                ("rate", "Hit fraction"),
                ("unreviewed", "Unreviewed"),
            ],
        ),
        "<h2>Reliability</h2><p>M7 status: "
        + esc(data["reliability"]["status"])
        + ". Undefined or missing kappa is not zero or perfect agreement. Comparisons use saved spreadsheet values; no rater independence is inferred.</p>",
        _html_table(
            data["reliability"]["codes"],
            [
                ("code", "Code"),
                ("n_pairs", "Pairs"),
                ("observed_agreement", "Observed agreement"),
                ("kappa", "Kappa"),
                ("comparison", "Sheet comparison"),
            ],
        ),
        '<h2>Order details</h2><div class="grid">',
    ]
    for i, c in enumerate(data["cases"], 1):
        parts.append(
            f'<div><h3>{esc(c["case_id"])}</h3><p>{esc(c["utility"] or "Utility unrecorded")} | {esc(c["order_date"] or "Date missing")}</p><img src="figures/orders/{i:02d}-{esc(c["case_id"])}.png" alt="ROE values for {esc(c["case_id"])}"><details><summary>Source record</summary><p>{esc(str(c["order_url"] or "URL missing"))}</p><p>{esc(str(c["sources_notes"] or "Notes missing"))}</p></details></div>'
        )
    parts += [
        "</div><h2>Coverage and methods</h2>",
        _html_table(
            data["feasibility"]["field_coverage"],
            [
                ("field", "Field"),
                ("entered", "Entered"),
                ("orders", "Orders"),
                ("share", "Coverage fraction"),
            ],
        ),
        "<details><summary>Excluded scatter intervals</summary>"
        + _html_table(
            data["scatter"]["excluded"],
            [("case_id", "Order"), ("missing_fields", "Missing fields")],
        )
        + "</details>",
        "<ul>" + "".join("<li>" + esc(m) + "</li>" for m in METHODS) + "</ul>",
        "<p>Input: "
        + esc(Path(data["source_path"]).name)
        + "<br>SHA-256: <code>"
        + data["source_sha256"]
        + "</code></p></main></html>",
    ]
    (destination / "report.html").write_text("\n".join(parts))


def _csv(path, rows, headers):
    def value(v):
        if isinstance(v, (list, dict)):
            v = json.dumps(v, ensure_ascii=False)
        if isinstance(v, str) and v.startswith(("=", "+", "-", "@", "\t", "\r")):
            return "'" + v
        return v

    with path.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({k: value(r.get(k)) for k in headers} for r in rows)


def _pdf_table(fig, bounds, title, rows, headers):
    ax = fig.add_axes(bounds)
    ax.axis("off")
    ax.set_title(title, loc="left", fontsize=10, fontweight="bold", pad=8)
    if not rows:
        ax.text(0, 0.9, "No observations recorded.", fontsize=9)
        return
    height = min(0.96, (len(rows) + 1) * 0.024 / bounds[3])
    table = ax.table(
        cellText=[[_display(v) for v in r] for r in rows],
        colLabels=headers,
        cellLoc="left",
        loc="upper left",
        bbox=[0, 0.96 - height, 1, height],
    )
    table.auto_set_font_size(False)
    table.set_fontsize(8)
    for (row, col), cell in table.get_celld().items():
        cell.set_edgecolor("#d5dfe5")
        cell.set_linewidth(0.35)
        if row == 0:
            cell.set_facecolor("#e8eef2")
            cell.set_text_props(weight="bold")


def _memo_pdf(data, destination):
    """Two landscape pages: descriptive figures, then feasibility and reliability."""
    meta = {
        "Title": "Anchored Returns - descriptive pilot memo",
        "Author": "Anchored Returns project",
        "CreationDate": None,
        "ModDate": None,
        "Subject": _banner(data),
    }
    with PdfPages(destination / "memo.pdf", metadata=meta) as pdf:
        fig = plt.figure(figsize=(11, 8.5))
        fig.text(
            0.04, 0.95, "Anchored Returns | Descriptive pilot memo", fontsize=19, weight="bold"
        )
        fig.text(0.04, 0.915, _banner(data), fontsize=10, color="#8b4c24")
        fig.text(
            0.04,
            0.879,
            f"Orders: {data['n_cases']} | Scatter: descriptive, n = {data['scatter']['n_plotted']} | Fully coded analysis passages: {data['coding']['n_complete_commission']}",
            fontsize=10,
        )
        for path, bounds in [
            ("number_lines.png", [0.015, 0.13, 0.57, 0.72]),
            ("rate_shock.png", [0.59, 0.46, 0.39, 0.39]),
            ("code_shares.png", [0.59, 0.11, 0.39, 0.34]),
        ]:
            ax = fig.add_axes(bounds)
            ax.imshow(plt.imread(destination / "figures" / path))
            ax.axis("off")
        fig.text(
            0.04,
            0.074,
            "Missing inputs are omitted, not replaced by zero. Scatter uses base ROE and matched prior/current dates.",
            fontsize=9,
        )
        fig.text(
            0.04,
            0.047,
            "Descriptive figures do not establish anchoring, causation or independent coding. Per-order values and sources: report.html.",
            fontsize=9,
        )
        fig.text(0.95, 0.025, "1 / 2", fontsize=8)
        pdf.savefig(fig)
        plt.close(fig)
        fig = plt.figure(figsize=(11, 8.5))
        fig.text(0.04, 0.95, "Feasibility, coverage and reliability", fontsize=18, weight="bold")
        fig.text(0.04, 0.915, _banner(data), fontsize=9, color="#8b4c24")
        f = data["feasibility"]
        rows = [
            [
                r["case_id"],
                r["minutes_spent"],
                f"{r['numeric_fields_entered']}/{r['numeric_fields_tracked']}",
                r["n_complete_commission"],
            ]
            for r in f["orders"]
        ]
        _pdf_table(
            fig,
            [0.04, 0.54, 0.50, 0.32],
            f"Order feasibility | median minutes: {_display(f['median_minutes'])}",
            rows,
            ["Order", "Minutes", "Fields", "Analysis"],
        )
        coverage = [
            [r["field"].replace("_", " "), f"{r['entered']}/{r['orders']}"]
            for r in f["field_coverage"]
        ]
        _pdf_table(
            fig,
            [0.04, 0.12, 0.50, 0.36],
            "Entered numeric fields (not source certification)",
            coverage,
            ["Field", "Coverage"],
        )
        s = _sensitivity_view(data)
        _pdf_table(
            fig,
            [0.59, 0.715, 0.37, 0.13],
            "Expected earnings sensitivity",
            [[r["variant"], r["market"], r["benchmark"]] for r in s],
            ["Treatment", "Market", "Benchmark"],
        )
        _pdf_table(
            fig,
            [0.59, 0.50, 0.37, 0.15],
            "M3 top-hit reviews",
            [
                [
                    r["reviewer_type"],
                    f"{r['hits']}/{r['reviewed']}",
                    None if r["rate"] is None else f"{r['rate']:.0%}",
                ]
                for r in f["m3_hit_rates"]
            ],
            ["Reviewer", "Hits / n", "Rate"],
        )
        _pdf_table(
            fig,
            [0.59, 0.18, 0.37, 0.26],
            "Reliability | " + data["reliability"]["status"],
            [[r["code"], r["n_pairs"], r["kappa"]] for r in data["reliability"]["codes"]],
            ["Code", "Pairs", "Kappa"],
        )
        fig.text(
            0.04,
            0.075,
            "Shares use fully coded Commission analysis. Bounds flag ambiguous methods. Missing ratings do not establish reliability.",
            fontsize=9,
        )
        fig.text(0.04, 0.05, "Workbook SHA-256: " + data["source_sha256"], fontsize=7)
        fig.text(
            0.04,
            0.025,
            "Full provenance, exclusions, denominators and methods are in report.html and analysis.json.",
            fontsize=8,
        )
        fig.text(0.95, 0.025, "2 / 2", fontsize=8)
        pdf.savefig(fig)
        plt.close(fig)


def _destination(path):
    p = Path(path).expanduser().absolute()
    _require(not p.exists(), "Output directory already exists; choose a fresh report directory.")
    _require(
        not any(v.is_symlink() for v in (p, *p.parents)), "Symlink output paths are not supported."
    )
    _require(
        not PROTECTED.intersection(p.parts), "Output is in a protected source/coding directory."
    )
    return p


def build_report(workbook, output_dir, *, scope="pilot", synthetic=False):
    """Write a new complete report bundle after validating all workbook inputs.

    Refuse existing/protected destinations; retain a completion audit only on success.
    """
    destination = _destination(output_dir)
    data = analyse_workbook(workbook, scope=scope, synthetic=synthetic)
    destination.mkdir(parents=True, exist_ok=False)
    try:
        figures = destination / "figures"
        (figures / "orders").mkdir(parents=True)
        tables = destination / "tables"
        tables.mkdir()
        with plt.rc_context(
            {
                "font.family": "DejaVu Sans",
                "font.size": 10,
                "svg.hashsalt": "anchored-returns-m8",
                "axes.unicode_minus": False,
                "text.parse_math": False,
            }
        ):
            for name, fig in [
                ("rate_shock", plot_rate_shock(data)),
                ("number_lines", plot_number_lines(data)),
                ("code_shares", plot_code_shares(data)),
            ]:
                fig.savefig(
                    figures / (name + ".png"), dpi=180, metadata={"Software": "Anchored Returns M8"}
                )
                fig.savefig(figures / (name + ".svg"), metadata={"Date": None})
                plt.close(fig)
            for i, c in enumerate(data["cases"], 1):
                fig = plot_number_lines(data, c)
                for suffix in ("png", "svg"):
                    fig.savefig(
                        figures / "orders" / f"{i:02d}-{c['case_id']}.{suffix}",
                        dpi=180,
                        metadata={"Date": None}
                        if suffix == "svg"
                        else {"Software": "Anchored Returns M8"},
                    )
                plt.close(fig)
            _memo_pdf(data, destination)
        _csv(
            tables / "number_lines.csv",
            data["cases"],
            [
                "case_id",
                "state",
                "order_date",
                "regime",
                *NUMBER_FIELDS,
                "perf_adjustment_bp",
                "order_url",
                "sources_notes",
                "source_row",
            ],
        )
        _csv(
            tables / "rate_shock.csv",
            data["scatter"]["points"],
            [
                "case_id",
                "state",
                "regime",
                "prior_order_date",
                "order_date",
                "delta_roe_pp",
                "delta_ust10_pp",
            ],
        )
        _csv(
            tables / "rate_shock_exclusions.csv",
            data["scatter"]["excluded"],
            ["case_id", "missing_fields"],
        )
        for name, rows, headers in [
            (
                "code_shares",
                data["code_shares"],
                ["group", "code", "count", "denominator", "share"],
            ),
            ("expected_earnings", data["sensitivity"], list(data["sensitivity"][0])),
            (
                "feasibility",
                data["feasibility"]["orders"],
                [
                    "case_id",
                    "state",
                    "regime",
                    "minutes_spent",
                    "numeric_fields_entered",
                    "numeric_fields_tracked",
                    "roe_section_recorded",
                    "n_passages",
                    "n_commission",
                    "n_complete_commission",
                ],
            ),
            (
                "field_coverage",
                data["feasibility"]["field_coverage"],
                ["field", "entered", "orders", "share"],
            ),
            (
                "m3_hit_rates",
                data["feasibility"]["m3_hit_rates"],
                ["reviewer_type", "hits", "reviewed", "rate", "unreviewed"],
            ),
            (
                "reliability",
                data["reliability"]["codes"],
                [
                    "code",
                    "n_pairs",
                    "observed_agreement",
                    "expected_agreement",
                    "kappa",
                    "sheet_kappa",
                    "comparison",
                ],
            ),
        ]:
            _csv(tables / (name + ".csv"), rows, headers)
        (destination / "analysis.json").write_text(
            json.dumps(data, indent=2, ensure_ascii=False, allow_nan=False) + "\n"
        )
        _write_html(data, destination)
        lines = [
            "# Anchored Returns: descriptive pilot memo",
            "",
            _banner(data),
            "",
            f"Populated orders: {data['n_cases']}. Comparable intervals: {data['scatter']['n_plotted']}. Fully coded Commission-analysis passages: {data['coding']['n_complete_commission']}.",
            "",
            f"Median recorded minutes: {_display(data['feasibility']['median_minutes'])}; timing coverage {data['feasibility']['minutes_recorded']}/{data['n_cases']}. M7 status: {data['reliability']['status']}.",
            "",
            "[Full tables, sources and order details](report.html) | [Two-page memo](memo.pdf)",
            "",
            "![ROE comparisons](figures/number_lines.png)",
            "",
            "![Rate shock](figures/rate_shock.png)",
            "",
            "![Code shares](figures/code_shares.png)",
            "",
            "## Expected earnings sensitivity",
            "",
        ]
        for r in _sensitivity_view(data):
            lines.append(
                f"- {r['variant']}: market {r['market']}; benchmark {r['benchmark']}; ambiguous methods {r['ambiguous']}."
            )
        lines += (
            ["", "## Methods and limitations", ""]
            + ["- " + m for m in METHODS]
            + ["", f"Input SHA-256: {data['source_sha256']}"]
        )
        (destination / "memo.md").write_text("\n".join(lines) + "\n")
        _require(
            _hash(workbook) == data["source_sha256"], "Input changed during report generation."
        )
        files = [
            dict(path=str(p.relative_to(destination)), sha256=_hash(p), bytes=p.stat().st_size)
            for p in sorted(destination.rglob("*"))
            if p.is_file()
        ]
        audit = dict(
            module="M8",
            version=VERSION,
            status="complete",
            dataset_kind=data["dataset_kind"],
            source_path=data["source_path"],
            source_sha256=data["source_sha256"],
            config=dict(scope=scope, synthetic=synthetic),
            n_cases=data["n_cases"],
            n_scatter=data["scatter"]["n_plotted"],
            files=files,
            workbook_validation_status=data["validation"]["status"],
            reliability_comparison_status=data["reliability"]["status"],
            limits=[
                "Completion means all requested output files were written, not that research data are complete.",
                "Figures and tables use only workbook inputs; source accuracy and sampling eligibility are not certified.",
            ],
        )
        (destination / "audit.json").write_text(json.dumps(audit, indent=2) + "\n")
        return audit
    except Exception:
        # This directory was created exclusively by this call; remove partial output only.
        shutil.rmtree(destination)
        raise


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("workbook", type=Path)
    p.add_argument("output_dir", type=Path)
    p.add_argument("--scope", choices=("pilot", "practice"), default="pilot")
    p.add_argument(
        "--synthetic",
        action="store_true",
        help="Mark every output as a software demonstration, not research data.",
    )
    args = p.parse_args(argv)
    try:
        r = build_report(args.workbook, args.output_dir, scope=args.scope, synthetic=args.synthetic)
    except (ReportInputError, OSError, ValueError) as e:
        print(f"error: {e}", file=sys.stderr)
        return 2
    print(
        json.dumps(
            dict(
                status=r["status"],
                dataset_kind=r["dataset_kind"],
                n_cases=r["n_cases"],
                n_scatter=r["n_scatter"],
                workbook_validation_status=r["workbook_validation_status"],
                reliability_comparison_status=r["reliability_comparison_status"],
                output_dir=str(args.output_dir.resolve()),
            ),
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
