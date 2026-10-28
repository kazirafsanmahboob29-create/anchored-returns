"""M7: read-only, per-code Cohen kappa and saved-workbook cross-checks.

No ratings are inferred or repaired. Saved formula values are compared, never
executed here; a matching cache is not evidence of independent raters.
"""

from __future__ import annotations

import argparse
from collections import Counter
from decimal import Decimal
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import sys
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.exceptions import InvalidFileException
import sklearn
from sklearn.metrics import cohen_kappa_score

from src.validate_sheet import SCHEMA, ID, _blank, _signature

SCHEMA_PATH = Path(__file__).with_name("reliability_schema_v012.json")
RELIABILITY_SCHEMA = json.loads(SCHEMA_PATH.read_text())
CODES = tuple(RELIABILITY_SCHEMA["codes"])
KAPPA_TOLERANCE = Decimal("0.001")


class ReliabilityInputError(ValueError):
    """Malformed workbook or incomplete/invalid paired ratings."""


class ReliabilityMismatchError(ValueError):
    """A saved workbook result disagrees with the recomputed result."""

    def __init__(self, report):
        self.report = report
        super().__init__("Saved reliability results disagree; see report issues.")


def _number(value):
    return type(value) in (int, float) and math.isfinite(value)


def _report_value(value):
    # Invalid dates/nonfinite caches must still yield a JSON-readable diagnostic.
    return value if value is None or type(value) in (str, bool) or _number(value) else str(value)


def binary_agreement(a, b):
    """Compute one unweighted binary kappa using scikit-learn.

    Also expose counts and marginal probabilities, with null for undefined kappa.
    """
    a, b = list(a), list(b)
    if len(a) != len(b) or any(not _number(x) or x not in (0, 1) for x in a + b):
        raise ReliabilityInputError(
            "Paired vectors must have equal lengths and numeric 0/1 values only."
        )
    counts = Counter(zip(a, b))
    n = len(a)
    yes_a, yes_b = sum(a), sum(b)
    observed = (counts[0, 0] + counts[1, 1]) / n if n else None
    expected = (yes_a * yes_b + (n - yes_a) * (n - yes_b)) / (n * n) if n else None
    # Single common category is mathematically 0/0. Do not replace it by 1.
    kappa = float(cohen_kappa_score(a, b, labels=[0, 1])) if n and expected != 1 else None
    if kappa is not None and not math.isfinite(kappa):
        raise ReliabilityInputError("Unexpected nonfinite scikit-learn kappa.")
    return dict(
        status="no_pairs" if not n else "undefined" if kappa is None else "defined",
        n_pairs=n,
        both_1=counts[1, 1],
        both_0=counts[0, 0],
        a_only=counts[1, 0],
        b_only=counts[0, 1],
        a_yes=int(yes_a),
        b_yes=int(yes_b),
        observed_agreement=observed,
        expected_agreement=expected,
        kappa=kappa,
        meets_0_70=None if kappa is None else kappa >= 0.7,
    )


def _require(condition, message):
    if not condition:
        raise ReliabilityInputError(message)


def _read_inputs(book):
    """Validate the fixed template and reference identities before reading ratings.

    Other research fields remain M6's responsibility; blank drafts are not coded.
    """
    for sheet, headers in [
        ("Cases", SCHEMA["case_headers"]),
        ("Passages", SCHEMA["passage_headers"]),
        ("Reliability", RELIABILITY_SCHEMA["headers"]),
    ]:
        _require(sheet in book, f"Missing worksheet {sheet}.")
        ws = book[sheet]
        _require(ws.max_row <= 10000 and ws.max_column <= 256, f"{sheet}: unsupported dimensions.")
        _require(
            [ws.cell(1, c).value for c in range(1, len(headers) + 1)] == headers,
            f"{sheet}: expected original v0.1.2 headers.",
        )
    ws = book["Reliability"]
    for row, code in enumerate(CODES, 2):
        _require(ws[f"F{row}"].value == code, f"Reliability!F{row}: unexpected summary code/order.")
        for origin, formula in RELIABILITY_SCHEMA["formulas"].items():
            target = origin[0] + str(row)
            expected = Translator(formula, origin=origin).translate_formula(target)
            _require(
                ws[target].value == expected, f"Reliability!{target}: formula differs from v0.1.2."
            )
    cases, passages = set(), set()
    for sheet, width, column, target, example in [
        ("Cases", 41, 1, cases, "EX-01"),
        ("Passages", 21, 1, passages, "EX-01-001"),
    ]:
        seen = set()
        for row in range(2, book[sheet].max_row + 1):
            cells = [book[sheet].cell(row, c) for c in range(1, width + 1)]
            values = [c.value for c in cells]
            identity = values[column - 1]
            if all(_blank(v) for v in values):
                continue
            if identity == example or (sheet == "Passages" and values[1] == "EX-01"):
                _require(
                    _signature(cells) == SCHEMA["example_signatures"][sheet],
                    f"{sheet} row {row}: modified example.",
                )
                continue
            if sheet == "Cases" and all(_blank(v) for v in values[1:]):
                continue  # ID-only placeholders do not create valid case references.
            _require(
                isinstance(identity, str) and ID.fullmatch(identity),
                f"{sheet}!A{row}: invalid identity.",
            )
            _require(identity not in seen, f"{sheet}!A{row}: duplicate identity {identity}.")
            _require(
                row <= (17 if sheet == "Cases" else 400),
                f"{sheet}!A{row}: outside template capacity.",
            )
            seen.add(identity)
            if sheet == "Passages":
                cid = values[1]
                _require(
                    cid in cases
                    and identity.startswith(cid + "-")
                    and len(identity) > len(cid) + 1,
                    f"Passages!B{row}: missing case or incorrect passage prefix.",
                )
            target.add(identity)
    groups = {code: [] for code in CODES}
    seen = set()
    for row in range(2, ws.max_row + 1):
        cells = [ws.cell(row, c) for c in range(1, 5)]
        values = [c.value for c in cells]
        if all(_blank(v) for v in values):
            continue
        _require(row <= 601, f"Reliability!A{row}: beyond formula coverage (row 601).")
        for cell in cells:
            _require(
                cell.data_type not in ("f", "e"),
                f"Reliability!{cell.coordinate}: formula/error in input.",
            )
        pid, code, a, b = values
        _require(
            isinstance(pid, str) and pid in passages,
            f"Reliability!A{row}: unknown or illustrative passage ID.",
        )
        _require(code in CODES, f"Reliability!B{row}: unknown code.")
        for cell in cells[2:]:
            _require(
                _number(cell.value) and cell.value in (0, 1),
                f"Reliability!{cell.coordinate}: paired rating must be numeric 0 or 1; no missing ratings are dropped.",
            )
        _require((pid, code) not in seen, f"Reliability!A{row}: duplicate passage/code pair.")
        seen.add((pid, code))
        groups[code].append((pid, int(a), int(b)))
    return groups


def _compare(saved, row, result, issues):
    """Compare saved summary cells with independent counts and sklearn kappa.

    Missing caches remain unavailable; numeric discrepancies become exceptions.
    """
    keys = (
        "n_pairs",
        "both_1",
        "both_0",
        "a_yes",
        "b_yes",
        "observed_agreement",
        "expected_agreement",
        "kappa",
        "meets_0_70",
    )
    missing, failed = [], []
    for col, key in zip("GHIJKLMNO", keys):
        cell = f"{col}{row}"
        value = saved[cell].value
        expected = result[key]
        if key == "meets_0_70":
            expected = None if expected is None else "Yes" if expected else "Rewrite"
        if value is None or value == "":
            if expected is not None and (result["n_pairs"] or expected != 0):
                missing.append(cell)
            continue
        if key == "kappa":
            result["sheet_kappa"] = _report_value(value)
        equal = False
        if expected is not None:
            if isinstance(expected, str):
                equal = value == expected
            elif _number(value):
                tolerance = (
                    KAPPA_TOLERANCE
                    if key == "kappa"
                    else Decimal("0")
                    if col in "GHIJK"
                    else Decimal("1e-12")
                )
                difference = abs(Decimal(str(value)) - Decimal(str(expected)))
                equal = difference <= tolerance
                if key == "kappa":
                    result["absolute_difference"] = float(difference)
        if not equal:
            failed.append(cell)
            issues.append(
                dict(
                    code="kappa_mismatch" if key == "kappa" else "summary_mismatch",
                    cell="Reliability!" + cell,
                    expected=expected,
                    saved=_report_value(value),
                )
            )
    if failed:
        return "mismatch"
    if missing:
        issues.append(
            dict(
                code="missing_cache",
                cells=["Reliability!" + c for c in missing],
                message="Recalculate and save a separate copy in a spreadsheet engine, then rerun.",
            )
        )
        return "missing_cache"
    return (
        "not_applicable"
        if not result["n_pairs"]
        else "undefined"
        if result["kappa"] is None
        else "matched"
    )


def reliability(path, *, compute_only=False):
    """Compute each code separately and compare the saved workbook summary.

    Raise ReliabilityMismatchError carrying the full report for disagreements.
    """
    path = Path(path).expanduser().resolve()
    book = cached = None
    try:
        content = path.read_bytes()
        with ZipFile(BytesIO(content)) as archive:
            _require(
                sum(i.file_size for i in archive.infolist()) <= 100 * 1024 * 1024,
                "Workbook exceeds 100 MiB uncompressed limit.",
            )
        book = load_workbook(BytesIO(content), data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(content), data_only=True, keep_links=False)
        groups = _read_inputs(book)
        results, issues = [], []
        for row, code in enumerate(CODES, 2):
            pairs = groups[code]
            r = dict(
                code=code,
                **binary_agreement([x[1] for x in pairs], [x[2] for x in pairs]),
                sheet_cell=f"Reliability!N{row}",
                sheet_kappa=None,
                absolute_difference=None,
            )
            r["comparison"] = (
                "not_requested" if compute_only else _compare(cached["Reliability"], row, r, issues)
            )
            results.append(r)
        total = sum(r["n_pairs"] for r in results)
        mismatched = any(r["comparison"] == "mismatch" for r in results)
        defined = sum(r["status"] == "defined" for r in results)
        compared = sum(r["comparison"] == "matched" for r in results)
        unavailable = any(
            r["n_pairs"] and (r["status"] != "defined" or r["comparison"] == "missing_cache")
            for r in results
        )
        status = (
            "failed"
            if mismatched
            else "no_data"
            if not total
            else "computed_only"
            if compute_only
            else "incomplete"
            if unavailable
            else "passed"
        )
        report = dict(
            module="M7",
            version="1.0",
            workbook_schema=RELIABILITY_SCHEMA["version"],
            sklearn_version=sklearn.__version__,
            source_path=str(path),
            source_sha256=hashlib.sha256(content).hexdigest(),
            schema_sha256=hashlib.sha256(SCHEMA_PATH.read_bytes()).hexdigest(),
            config=dict(compute_only=compute_only, kappa_tolerance=str(KAPPA_TOLERANCE)),
            status=status,
            total_pairs=total,
            unique_passages=len({p[0] for pairs in groups.values() for p in pairs}),
            populated_codes=sum(bool(r["n_pairs"]) for r in results),
            defined_codes=defined,
            compared_codes=compared,
            codes=results,
            issues=issues,
            limits=[
                "Checks arithmetic consistency, not rater identity, independence or codebook validity.",
                "Saved formula results are inspected; this tool neither runs Excel nor saves the workbook.",
                "No-pair codes and undefined kappa are not perfect agreement; no pooled multi-label kappa is computed.",
                "The 0.70 flag is the supplied project target, not a significance test or a general quality guarantee.",
                "Run M6 separately for full workbook validation.",
            ],
        )
        if mismatched:
            raise ReliabilityMismatchError(report)
        return report
    except (ReliabilityInputError, ReliabilityMismatchError):
        raise
    except (
        OSError,
        BadZipFile,
        InvalidFileException,
        ValueError,
        KeyError,
        TypeError,
        SyntaxError,
    ) as error:
        raise ReliabilityInputError(f"Cannot read reliability workbook: {error}") from error
    finally:
        if book is not None:
            book.close()
        if cached is not None:
            cached.close()


def format_report(report):
    lines = [
        f"M7 {report['status']}: {report['total_pairs']} pairs; {report['compared_codes']} codes cross-checked.",
        "Code | Pairs | Observed agreement | Kappa | Saved kappa | Comparison",
    ]
    for r in report["codes"]:
        f = lambda value: "unavailable" if value is None else str(value)
        lines.append(
            f"{r['code']} | {r['n_pairs']} | {f(r['observed_agreement'])} | {f(r['kappa'])} | {f(r['sheet_kappa'])} | {r['comparison']}"
        )
    lines.extend(json.dumps(i, ensure_ascii=False) for i in report["issues"])
    lines.append("Arithmetic agreement does not establish independent coding reliability.")
    return "\n".join(lines) + "\n"


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--compute-only",
        action="store_true",
        help="Compute ratings without claiming a workbook cross-check.",
    )
    parser.add_argument("--format", choices=("text", "json"), default="text")
    args = parser.parse_args(argv)
    try:
        report = reliability(args.workbook, compute_only=args.compute_only)
    except ReliabilityMismatchError as error:
        report = error.report
    except ReliabilityInputError as error:
        print(f"error: {error}", file=sys.stderr)
        return 2
    print(
        json.dumps(report, indent=2, ensure_ascii=False, allow_nan=False)
        if args.format == "json"
        else format_report(report),
        end="\n" if args.format == "json" else "",
    )
    return 0 if report["status"] in ("passed", "computed_only") else 1


if __name__ == "__main__":
    raise SystemExit(main())
